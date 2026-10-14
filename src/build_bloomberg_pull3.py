"""
Third Bloomberg data-pull workbook: the blade supplier (Howmet), supply-chain links in dollars, and segment tables.

    python src/build_bloomberg_pull3.py      ->  data/bloomberg/GEV_bloomberg_pull_3.xlsx

Research data pull only, not the model the team submits. The Links and Segments sheets use field names that
have not been tested on a terminal; the Manual sheet holds the screen exports that are certain to work.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from build_bloomberg_pull import BOLD, FILL, GEV, HEAD, NOTE, bdp, header, widths

OUT = Path(__file__).resolve().parent.parent / "data" / "bloomberg" / "GEV_bloomberg_pull_3.xlsx"
HWM = "HWM US Equity"

OUTLOOK = [("Howmet", HWM), ("ATI", "ATI US Equity"), ("Carpenter Technology", "CRS US Equity"), ("Woodward", "WWD US Equity"),
           ("GE Aerospace", "GE US Equity"), ("Siemens Energy", "ENR GY Equity"), ("Mitsubishi Heavy", "7011 JT Equity")]
TO_GEV = ["HWM US Equity", "WWD US Equity", "ATI US Equity", "CRS US Equity", "5631 JT Equity", "TKR US Equity", "5401 JP Equity",
          "TKA GR Equity", "ACA US Equity", "112610 KS Equity", "GE US Equity"]
HWM_TO = ["GEV US Equity", "ENR GY Equity", "7011 JP Equity", "GE US Equity", "RTX US Equity", "SAF FP Equity", "RR/ LN Equity",
          "BA US Equity", "AIR FP Equity", "BKR US Equity"]
GEV_TO = ["DUK US Equity", "SO US Equity", "ETR US Equity", "NEE US Equity", "D US Equity", "AEP US Equity", "XEL US Equity", "VST US Equity",
          "CEG US Equity", "NRG US Equity", "TLN US Equity", "IBE SM Equity", "015760 KS Equity", "NG/ LN Equity", "RWE GR Equity",
          "ENEL IM Equity", "ENGI FP Equity", "9503 JP Equity"]


def rel(sec, related, side):
    """Dollar value of the link between two companies. side: C = related company is a customer, S = a supplier."""
    return ('=BDP("%s","RELATIONSHIP_AMOUNT","RELATIONSHIP_OVERRIDE=%s","QUANTIFIED_OVERRIDE=Y","EQY_FUND_CRNCY=USD",'
            '"RELATED_COMPANY_OVERRIDE=%s")' % (sec, side, related))


def readme(wb):
    ws = wb.active
    ws.title = "README"
    lines = [
        ("GE Vernova Bloomberg data pull, part 3 (short)", BOLD),
        ("", None),
        ("1. Open in Excel on a Bloomberg terminal PC, logged in. Wait about one minute until no cell says 'Requesting Data'.", None),
        ("2. Save As a new name and send the saved file back.", None),
        ("3. Then do the two screen exports on the Manual sheet. They take about five minutes each and are the sure way to get the data.", None),
        ("", None),
        ("The Outlook sheet uses field names that worked before. The Links and Segments sheets are tries: #N/A there is fine, leave it.", None),
        ("Bloomberg data is licensed for academic use. Keep the filled file within the team and cite as Bloomberg with the date.", None),
    ]
    for i, (text, font) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        if font:
            c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 130})


def outlook(wb):
    ws = wb.create_sheet("Outlook")
    ws.cell(row=1, column=1, value="What analysts expect each supplier and turbine maker to sell in the next four quarters. 1FQ is the next unreported quarter.").font = NOTE
    quarters = ("1FQ", "2FQ", "3FQ", "4FQ")
    header(ws, 2, ["Company", "Ticker", "Currency", "Next report date"] + ["Revenue " + q for q in quarters] + ["Estimates, 1FQ", "EPS 1FQ", "EBITDA 1FQ", "EBITDA 2FQ"])
    for r, (name, tic) in enumerate(OUTLOOK, 3):
        row = [name, tic, bdp(tic, "CRNCY"), bdp(tic, "EXPECTED_REPORT_DT")] + [bdp(tic, "BEST_SALES", q) for q in quarters]
        row += [bdp(tic, "BEST_SALES_NUMEST", "1FQ"), bdp(tic, "BEST_EPS", "1FQ"), bdp(tic, "BEST_EBITDA", "1FQ"), bdp(tic, "BEST_EBITDA", "2FQ")]
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    widths(ws, {"A": 22, "B": 16, "C": 10, "D": 16})
    for c in range(5, 13):
        ws.column_dimensions[get_column_letter(c)].width = 14


def links(wb):
    ws = wb.create_sheet("Links")
    ws.cell(row=1, column=1, value="Supply-chain links in dollars (Bloomberg's estimates, as shown on SPLC). Untested field names: #N/A is fine.").font = NOTE
    blocks = [("What each supplier sells to GE Vernova", TO_GEV, lambda t: (rel(t, GEV, "C"), rel(GEV, t, "S")), ("Supplier's view", "GE Vernova's view")),
              ("What Howmet sells to each engine and turbine maker", HWM_TO, lambda t: (rel(HWM, t, "C"), rel(t, HWM, "S")), ("Howmet's view", "Customer's view")),
              ("What GE Vernova sells to each customer", GEV_TO, lambda t: (rel(GEV, t, "C"), rel(t, GEV, "S")), ("GE Vernova's view", "Customer's view"))]
    row = 3
    for title, tickers, make, views in blocks:
        ws.cell(row=row, column=1, value=title).font = BOLD
        header(ws, row + 1, ["Ticker", "Name", views[0], views[1], "Company revenue, last year", "Currency"])
        for k, t in enumerate(tickers, row + 2):
            a, b = make(t)
            for c, v in enumerate([t, bdp(t, "NAME"), a, b, bdp(t, "SALES_REV_TURN"), bdp(t, "CRNCY")], 1):
                ws.cell(row=k, column=c, value=v)
        row += len(tickers) + 4
    widths(ws, {"A": 18, "B": 30, "C": 20, "D": 20, "E": 24, "F": 10})


def segments(wb):
    ws = wb.create_sheet("Segments")
    ws.cell(row=1, column=1, value="Revenue by business line, as reported. Each table spills down and to the right. Untested field names: #N/A is fine.").font = NOTE
    col = 1
    for name, tic in (("Howmet", HWM), ("Siemens Energy", "ENR GY Equity"), ("Mitsubishi Heavy", "7011 JT Equity"), ("GE Vernova", GEV)):
        for label, extra in (("plain", ""), ("by product", ',"PRODUCT_GEO_OVERRIDE=P"')):
            c = ws.cell(row=2, column=col, value="%s, %s  [PG_REVENUE]" % (name, label))
            c.font, c.fill = HEAD, FILL
            ws.cell(row=3, column=col, value='=BDS("%s","PG_REVENUE"%s)' % (tic, extra))
            for k in range(6):
                ws.column_dimensions[get_column_letter(col + k)].width = 18
            col += 7


def manual(wb):
    ws = wb.create_sheet("Manual")
    header(ws, 1, ["#", "Type this", "Then", "Save as", "Why"])
    rows = [
        ("1", "HWM US Equity EEO <GO>", "Click the Company Specific tab. Set periodicity to Quarterly. Actions > Export to Excel. If there is a row for gas turbines or industrial gas turbines, that is the one we want.",
         "31_hwm_company_specific.xlsx", "The blade maker's gas turbine sales and what analysts expect next."),
        ("2", "BI <GO>", "Type 'gas turbines' in the search box, or open the Electrical Equipment industry page. Look for a table of gas turbine orders or market share by maker. Export it, or take a screenshot if it will not export.",
         "40_bi_gas_turbines.xlsx or .png", "An industry number by maker that can be cited."),
        ("3 (only if 2 finds nothing)", "GEV US Equity BI <GO>", "Open the company's Bloomberg Intelligence page. Export any table on gas turbine orders, backlog or pricing.",
         "41_bi_gev.xlsx or .png", "Same purpose as 2."),
    ]
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 14, "B": 26, "C": 80, "D": 32, "E": 60})


def main():
    wb = Workbook()
    readme(wb)
    outlook(wb)
    links(wb)
    segments(wb)
    manual(wb)
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    n = sum(1 for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    print("wrote %s: %d sheets, %d Bloomberg formulas" % (OUT, len(wb.sheetnames), n))


if __name__ == "__main__":
    main()
