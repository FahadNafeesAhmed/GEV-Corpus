"""Add the team's research files, the Bloomberg pulls and the reference model to the corpus, so one search covers everything.

    python src/compile_research.py            add or update, then index the changed text
    python src/compile_research.py --dry      list what would be added
    python src/compile_research.py --web      also fetch the web pages cited on 2 October 2026

What it writes to sources and chunks (chunks.form is the kind in capitals):
    research    the team's own research notes                         tier internal
    bloomberg   Bloomberg pulls rendered as text                      tier licensed (academic licence: keep inside the team)
    model       the reference model built during research             tier internal (not the team's submission)
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import warnings
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import load_workbook

from build_index import chunk_key, window
from common import ROOT, connect, html_to_text, load_config, log
from ingest_web import S, clean, pdf_text, store

warnings.filterwarnings("ignore")
DOCS = ROOT / "docs"
WINNERS = ROOT.parent / "pitch-winners-corpus"
BBG = ROOT / "data" / "bloomberg"
THESIS_SRC = ROOT / "scripts" / "thesis_pdf"
THESIS_MD = DOCS / "research" / "multi_agent" / "14_full_thesis" / "GEV_FULL_THESIS.md"
WEB = [
    ("https://www.howmet.com/press-release/2026-08-06/howmet-aerospace-reports-second-quarter-2026-results/", "web", "primary"),
    ("https://www.fool.com/earnings/call-transcripts/2026/08/13/howmet-aerospace-hwm-q2-2026-earnings-call-transcript/", "web", None),
    ("https://finance.yahoo.com/markets/stocks/articles/hwm-q2-earnings-beat-estimates-164300840.html", "news", None),
    ("https://www.thestreet.com/investing/stocks/gev-ge-vernova-ceo-strazik-says-200b-backlog-target-on-track-very-early-2027", "news", None),
    ("https://www.power-eng.com/nuclear/smrs/tennessee-clinch-river-bwrx-300-smr-nrc-construction-permit/", "news", None),
    ("https://www.nrg.com/assets/documents/insights/white-papers-and-webinars/pjm-bra-results-july2026.pdf", "web", None),
    ("https://www.gevernova.com/news/press-releases/ge-vernova-declares-fourth-quarter-2026-dividend", "web", "primary"),
    ("https://www.utilitydive.com/news/ge-vernova-gas-turbine-backlog-climbs-to-116-gw/826039/", "news", None),
    ("https://www.siemens-energy.com/global/en/home/press-releases/earnings-release-q3-fy-2026.html", "web", "primary"),
    ("https://www.investing.com/news/analyst-ratings/evercore-isi-reiterates-ge-vernova-stock-rating-on-raised-estimates-93CH-4926745", "news", None),
]


def fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return ("%.6g" % v) if abs(v) < 1e6 else ("%.0f" % v)
    return re.sub(r"\s+", " ", str(v)).strip()


def sheet_text(ws, max_rows: int = 600) -> str:
    rows = list(ws.iter_rows(max_row=max_rows + 5, values_only=True))
    head = [fmt(v).lower() for v in (rows[2] if len(rows) > 2 else [])]
    if head[:2] == ["date", "value"]:  # one Bloomberg history series every three columns
        out = []
        for c in range(0, len(head), 3):
            title = fmt(rows[1][c]) if c < len(rows[1]) else ""
            pts = ["%s %s" % (fmt(r[c]), fmt(r[c + 1])) for r in rows[3:] if c + 1 < len(r) and r[c] is not None and r[c + 1] is not None]
            if title and pts:
                out.append("%s\n%s" % (title, "\n".join(pts)))
        return "\n\n".join(out)
    lines = []
    for r in rows[:max_rows]:
        cells = [fmt(v) for v in r]
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def workbook_docs(path: Path, kind: str, tier: str, label: str, published: str):
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as e:  # unreadable export
        log("  skip %s: %s" % (path.name, e))
        return
    for ws in wb:
        body = sheet_text(ws)
        if len(body) < 200:
            continue
        rel = path.relative_to(ROOT).as_posix()
        yield dict(url="team://gev-corpus/%s#%s" % (rel, ws.title), kind=kind, tier=tier, angle=kind, published=published,
                   title="%s: %s, sheet %s" % (label, path.name, ws.title), text="%s\nFile: %s. Sheet: %s.\n\n%s" % (label, rel, ws.title, body))


def thesis_markdown() -> str:
    """The team thesis file as plain Markdown, from the same content the PDF is built from."""
    sys.path.insert(0, str(THESIS_SRC))
    import content_a, content_b, content_c
    out = ["# GE Vernova (GEV): the thesis file for the team", "",
           "_AI-assisted research notes, to be disclosed under rule 10. Not submission text. Bloomberg figures are licensed for academic use._", ""]
    for title, blocks in content_a.S + content_b.S + content_c.S + content_c.S_END:
        out += ["## " + title, ""]
        for b in blocks:
            if b[0] == "note":
                out += ["> Side note. " + b[1], ""]
            elif b[0] in ("md", "wide"):
                out += [b[1].strip(), ""]
            else:
                out += ["_Figure: %s._ %s" % (b[2], ("Source: %s." % b[3]) if b[3] else ""), ""]
    out += ["## 16. References", ""]
    out += ["- **%s** %s. %s" % (k, t, u or "") for k, t, u in content_c.REFS]
    return "\n".join(out) + "\n"


def docs():
    if THESIS_SRC.exists():
        THESIS_MD.parent.mkdir(parents=True, exist_ok=True)
        THESIS_MD.write_text(thesis_markdown(), encoding="utf-8")
    files = [(p, "gev-corpus", ROOT) for p in sorted(DOCS.rglob("*")) if p.suffix.lower() in (".md", ".csv") or p.name == "GEV_THESIS_DOSSIER.html"]
    files += [(p, "pitch-winners-corpus", WINNERS) for p in sorted(WINNERS.rglob("*.md"))] if WINNERS.exists() else []
    for p, root_name, root in files:
        raw = p.read_bytes()
        text = html_to_text(raw) if p.suffix.lower() == ".html" else raw.decode("utf-8", "replace")
        rel = p.relative_to(root).as_posix()
        m = re.search(r"^#\s+(.+)$", text, re.M)
        angle = "pitch_winners" if root is WINNERS else (p.parent.name if p.parent != DOCS else "docs")
        yield dict(url="team://%s/%s" % (root_name, rel), kind="research", tier="internal", angle=angle,
                   published=dt.date.fromtimestamp(p.stat().st_mtime).isoformat(), title=(m.group(1).strip() if m else p.stem)[:200] + " (" + rel + ")", text=text)
    lic = "Bloomberg data pulled on %s under an academic licence. Keep inside the team; cite as 'Bloomberg, %s'"
    for path, day in [(BBG / "GEV_bloomberg_pull_filled_2026-10-01.xlsx", "2026-10-01"), (BBG / "GEV_bloomberg_pull_2_filled_2026-10-02.xlsx", "2026-10-02")] + \
            [(p, "2026-10-01") for p in sorted((BBG / "GEV_Pull_2026-10-01").glob("*.xls*")) if p.name != "GEV_bloomberg_pull.xlsx"]:
        if path.exists():
            yield from workbook_docs(path, "bloomberg", "licensed", lic % (day, day), day)
    for p in (BBG / "street_segments_annual_2026-10-01.csv", BBG / "GEV_Pull_2026-10-01" / "segment_quarterly.json"):
        if p.exists():
            rel = p.relative_to(ROOT).as_posix()
            yield dict(url="team://gev-corpus/" + rel, kind="bloomberg", tier="licensed", angle="bloomberg", published="2026-10-01",
                       title="Bloomberg segment consensus: " + p.name, text=(lic % ("2026-10-01", "2026-10-01")) + "\nFile: %s\n\n%s" % (rel, p.read_text(encoding="utf-8", errors="replace")))
    model = ROOT / "data" / "model" / "GEV_model.xlsx"
    if model.exists():
        yield from workbook_docs(model, "model", "internal", "Reference model built during research (an answer key, not the team's submission)", "2026-10-02")


def fetch(item):
    url, kind, tier = item
    try:
        r = S.get(url, timeout=20, allow_redirects=True)
        if not r.ok:
            return None
        if "pdf" in r.headers.get("content-type", ""):
            return dict(url=url, kind=kind, tier=tier, angle="checks_2026-10-02", published="", title=url.rsplit("/", 1)[-1], text=pdf_text(r.content))
        m = re.search(r"<title[^>]*>(.*?)</title>", r.text, re.I | re.S)
        return dict(url=url, kind=kind, tier=tier, angle="checks_2026-10-02", published="", title=(m.group(1).strip() if m else url), text=clean(r.text))
    except Exception:
        return None


def index_source(con, sid: int, d: dict, size: int, overlap: int) -> int:
    """Re-chunk one source and keep the full-text index in step (delete old entries, add new ones)."""
    acc = "src:%d" % sid
    cols = "id, text, section, form, filing_date, accession, ticker"
    for row in con.execute("SELECT %s FROM chunks WHERE accession=?" % cols, (acc,)).fetchall():
        con.execute("INSERT INTO chunks_fts(chunks_fts, rowid, text, section, form, filing_date, accession, ticker) VALUES('delete',?,?,?,?,?,?,?)", row)
        con.execute("DELETE FROM embeddings WHERE chunk_id=?", (row[0],))
    con.execute("DELETE FROM chunks WHERE accession=?", (acc,))
    header = (d["title"] + "\n\n") if d["title"] else ""
    n = 0
    for i, piece in enumerate(window(header + d["text"], size, overlap)):
        vals = (acc, "GEV", d["kind"].upper(), (d["published"] or "")[:10], d["angle"], i, piece, d["tier"] or "commentary", d["kind"], chunk_key(acc, i, piece))
        cur = con.execute("INSERT INTO chunks (accession, ticker, form, filing_date, section, ord, text, tier, kind, chunk_key) VALUES (?,?,?,?,?,?,?,?,?,?)", vals)
        con.execute("INSERT INTO chunks_fts(rowid, text, section, form, filing_date, accession, ticker) VALUES (?,?,?,?,?,?,?)",
                    (cur.lastrowid, piece, d["angle"], d["kind"].upper(), (d["published"] or "")[:10], acc, "GEV"))
        n += 1
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry", action="store_true")
    ap.add_argument("--web", action="store_true")
    a = ap.parse_args()
    items = list(docs())
    if a.web:
        with ThreadPoolExecutor(6) as ex:
            got = [g for g in ex.map(fetch, WEB) if g and len(g["text"]) >= 300]
        log("web pages fetched: %d of %d" % (len(got), len(WEB)))
        items += got
    by_kind: dict = {}
    for d in items:
        k = by_kind.setdefault(d["kind"], [0, 0])
        k[0] += 1
        k[1] += len(d["text"])
    log("to compile: " + json.dumps({k: {"documents": v[0], "characters": v[1]} for k, v in by_kind.items()}))
    if a.dry:
        return
    con = connect()
    cfg = load_config()
    changed = chunks = 0
    for d in items:
        new = store(con, url=d["url"], kind=d["kind"], angle=d["angle"], title=d["title"], published=d["published"], text=d["text"], tier=d["tier"])
        row = con.execute("SELECT id FROM sources WHERE url=?", (d["url"],)).fetchone()
        if row is None:
            continue  # shorter than the corpus minimum of 300 characters
        has = con.execute("SELECT 1 FROM chunks WHERE accession=? LIMIT 1", ("src:%d" % row[0],)).fetchone()
        if new or not has:
            chunks += index_source(con, row[0], d, cfg["chunk_chars"], cfg["chunk_overlap"])
            changed += 1
    con.commit()
    log("sources added or updated: %d | chunks written: %d" % (changed, chunks))
    for kind, n, c in con.execute("SELECT kind, COUNT(DISTINCT accession), COUNT(*) FROM chunks WHERE kind IN ('research','bloomberg','model') GROUP BY kind"):
        log("  in corpus now: %-9s %4d documents, %6d chunks" % (kind, n, c))


if __name__ == "__main__":
    main()
