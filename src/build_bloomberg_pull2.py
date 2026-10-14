"""
Second Bloomberg data-pull workbook for GE Vernova: suppliers, customers, options, estimate detail.

Same idea as build_bloomberg_pull.py: only Bloomberg Excel add-in formulas, filled in on a terminal PC.

    python src/build_bloomberg_pull2.py      ->  data/bloomberg/GEV_bloomberg_pull_2.xlsx

This is a research data pull, not the model the team submits.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from build_bloomberg_pull import BOLD, FILL, GEV, HEAD, HYPERSCALERS, NOTE, bdh, bdp, header, history_sheet, widths

OUT = Path(__file__).resolve().parent.parent / "data" / "bloomberg" / "GEV_bloomberg_pull_2.xlsx"

SUPPLY = [("Howmet", "HWM US Equity", "turbine blades and vanes"), ("ATI", "ATI US Equity", "nickel alloys, forgings"),
          ("Carpenter Technology", "CRS US Equity", "superalloys"), ("Japan Steel Works", "5631 JT Equity", "rotor forgings; FY ends March"),
          ("Woodward", "WWD US Equity", "turbine fuel and control systems; FY ends September"), ("GE Aerospace", "GE US Equity", "buying the castings maker CPP"),
          ("Siemens Energy", "ENR GY Equity", "turbine maker; FY ends September"), ("Mitsubishi Heavy", "7011 JT Equity", "turbine maker; FY ends March"),
          ("Doosan Enerbility", "034020 KS Equity", "new turbine maker"), ("Baker Hughes", "BKR US Equity", "smaller turbines"),
          ("Caterpillar", "CAT US Equity", "Solar Turbines, engines"), ("Wartsila", "WRT1V FH Equity", "engines for on-site power"),
          ("Bloom Energy", "BE US Equity", "fuel cells for data centres"), ("FTAI Aviation", "FTAI US Equity", "jet engines converted to power turbines")]
CUSTOMERS = ["DUK US Equity", "SO US Equity", "ETR US Equity", "CEG US Equity", "NRG US Equity", "NEE US Equity", "D US Equity", "AEP US Equity",
             "XEL US Equity", "VST US Equity", "TLN US Equity", "PPL US Equity", "DTE US Equity", "AEE US Equity", "CNP US Equity", "WMB US Equity",
             "IBE SM Equity", "015760 KS Equity", "NG/ LN Equity", "RWE GR Equity", "EBK GR Equity", "ENEL IM Equity", "9963 TT Equity", "ENGI FP Equity",
             "ELI BB Equity", "PWGR IN Equity", "9503 JP Equity", "PPC GA Equity"]
FY = ("1FY", "2FY", "3FY")


def grid(wb, title, rows, fields, note):
    """rows: (name, ticker, note). fields: (label, field, period or None)."""
    ws = wb.create_sheet(title)
    ws.cell(row=1, column=1, value=note).font = NOTE
    header(ws, 2, ["Company", "Ticker", "Note"] + [f[0] for f in fields])
    header(ws, 3, ["Bloomberg field", "", ""] + [f[1] + (" [" + f[2] + "]" if f[2] else "") for f in fields])
    for r, (name, tic, nt) in enumerate(rows, 4):
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=tic)
        ws.cell(row=r, column=3, value=nt).font = NOTE
        for j, (_, field, per) in enumerate(fields):
            ws.cell(row=r, column=4 + j, value=bdp(tic, field, per))
    widths(ws, {"A": 22, "B": 18, "C": 34})
    for j in range(len(fields)):
        ws.column_dimensions[get_column_letter(4 + j)].width = 15
    ws.freeze_panes = "D4"


def bdh_q(sec, field, start):
    y, m, d = start
    return '=BDH("%s","%s",DATE(%d,%d,%d),TODAY(),"Per=Q","Days=A","Fill=P")' % (sec, field, y, m, d)


def readme(wb):
    ws = wb.active
    ws.title = "README"
    lines = [
        ("GE Vernova Bloomberg data pull, part 2", BOLD),
        ("", None),
        ("1. Open in Excel on a Bloomberg terminal PC, logged in, with the Bloomberg tab showing.", None),
        ("2. Wait until no cell says 'Requesting Data' (about two to three minutes). Use Refresh Workbook on the Bloomberg tab if needed.", None),
        ("3. Save As a new name with today's date and send the saved file back.", None),
        ("4. Then do the five screen exports on the Manual sheet. Bloomberg does not give those to Excel formulas.", None),
        ("", None),
        ("#N/A Invalid Field or #N/A N/A in a cell is fine. Leave it; it tells us that item is not available.", None),
        ("Bloomberg data is licensed for academic use. Keep the filled file within the team and cite as Bloomberg with the date.", None),
    ]
    for i, (text, font) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        if font:
            c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 130})


def suppliers(wb):
    fields = [("Name", "NAME", None), ("Currency", "CRNCY", None), ("Last price", "PX_LAST", None), ("Market cap (M)", "CUR_MKT_CAP", None),
              ("Next report date", "EXPECTED_REPORT_DT", None), ("Price change YTD %", "CHG_PCT_YTD", None), ("Price change 3M %", "CHG_PCT_3M", None),
              ("Revenue growth, trailing %", "SALES_GROWTH", None)]
    fields += [("Revenue " + p, "BEST_SALES", p) for p in FY] + [("EBITDA " + p, "BEST_EBITDA", p) for p in FY]
    fields += [("Fwd EV/EBITDA (NTM)", "BEST_EV_TO_BEST_EBITDA", "1BF"), ("Buy ratings", "TOT_BUY_REC", None), ("Avg price target", "BEST_TARGET_PRICE", None)]
    grid(wb, "Suppliers", SUPPLY, fields, "Blade, alloy and forging suppliers, the other turbine makers, and new sources of on-site power. 1FY is each company's own next unreported year.")

    series = []
    for name, tic, _ in SUPPLY[:8]:
        series.append(("%s quarterly revenue" % name, bdh_q(tic, "SALES_REV_TURN", (2023, 1, 1))))
    for name, tic, _ in SUPPLY[:3]:
        series.append(("%s revenue consensus 2FY (weekly)" % name, bdh(tic, "BEST_SALES", (2025, 1, 3), "2FY")))
        series.append(("%s EBITDA consensus 2FY (weekly)" % name, bdh(tic, "BEST_EBITDA", (2025, 1, 3), "2FY")))
    history_sheet(wb, "SupplierHistory", series, "Quarterly reported revenue since 2023, and whether supplier estimates are being raised.")


def customers(wb):
    fields = [("Name", "NAME", None), ("Currency", "CRNCY", None)] + [("Capex " + p, "BEST_CAPEX", p) for p in FY] + \
             [("Revenue 1FY", "BEST_SALES", "1FY"), ("Next report date", "EXPECTED_REPORT_DT", None)]
    grid(wb, "Customers", [("", t, "") for t in CUSTOMERS], fields, "Spending plans of GE Vernova's utility and power-producer customers. Capex is shown as a negative number by Bloomberg.")


def options(wb):
    ws = wb.create_sheet("Options")
    ws.cell(row=1, column=1, value="GE Vernova implied volatility: term structure and skew. Values are annualised %.").font = NOTE
    header(ws, 2, ["Item", "Bloomberg field", "Value"])
    items = [("%s, at the money" % t, "%s_IMPVOL_100.0%%MNY_DF" % c) for t, c in
             (("30 days", "30DAY"), ("60 days", "60DAY"), ("3 months", "3MTH"), ("6 months", "6MTH"), ("12 months", "12MTH"), ("18 months", "18MTH"), ("24 months", "24MTH"))]
    for t, c in (("30 days", "30DAY"), ("3 months", "3MTH"), ("12 months", "12MTH")):
        items += [("%s, 90%% strike (downside)" % t, "%s_IMPVOL_90.0%%MNY_DF" % c), ("%s, 110%% strike (upside)" % t, "%s_IMPVOL_110.0%%MNY_DF" % c)]
    items += [("3-month call implied vol (older field)", "3MO_CALL_IMP_VOL"), ("3-month put implied vol (older field)", "3MO_PUT_IMP_VOL"),
              ("Put/call open interest ratio", "PUT_CALL_OPEN_INTEREST_RATIO"), ("Put/call volume ratio", "PUT_CALL_VOLUME_RATIO_CUR_DAY"),
              ("Realised vol, 30 days", "VOLATILITY_30D"), ("Realised vol, 90 days", "VOLATILITY_90D"), ("Realised vol, 260 days", "VOLATILITY_260D"),
              ("Last price", "PX_LAST"), ("Next report date", "EXPECTED_REPORT_DT")]
    for r, (label, field) in enumerate(items, 3):
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=field).font = NOTE
        ws.cell(row=r, column=3, value=bdp(GEV, field))
    widths(ws, {"A": 40, "B": 34, "C": 14})
    series = [("GEV 3-month implied vol (weekly)", bdh(GEV, "3MTH_IMPVOL_100.0%MNY_DF", (2024, 4, 5))),
              ("GEV 12-month implied vol (weekly)", bdh(GEV, "12MTH_IMPVOL_100.0%MNY_DF", (2024, 4, 5)))]
    history_sheet(wb, "OptionsHistory", series, "Implied volatility since the spin.")


def estimates(wb):
    ws = wb.create_sheet("Estimates")
    ws.cell(row=1, column=1, value="GE Vernova estimate detail. Years: 1FY = 2026 ... 4FY = 2029. Quarters: 1FQ is the next unreported quarter (Q3 2026).").font = NOTE
    years = ("1FY", "2FY", "3FY", "4FY")
    metrics = [("Revenue mean", "BEST_SALES"), ("Revenue median", "BEST_SALES_MEDIAN"), ("Revenue high", "BEST_SALES_HI"), ("Revenue low", "BEST_SALES_LO"),
               ("Revenue std dev", "BEST_SALES_STDDEV"), ("Revenue number of estimates", "BEST_SALES_NUMEST"),
               ("Revenue change, 4 weeks", "BEST_SALES_4WK_CHG"), ("Revenue change, 3 months", "BEST_SALES_3MO_CHG"),
               ("EBITDA mean", "BEST_EBITDA"), ("EBITDA median", "BEST_EBITDA_MEDIAN"), ("EBITDA std dev", "BEST_EBITDA_STDDEV"),
               ("EBITDA number of estimates", "BEST_EBITDA_NUMEST"), ("EBITDA change, 4 weeks", "BEST_EBITDA_4WK_CHG"), ("EBITDA change, 3 months", "BEST_EBITDA_3MO_CHG"),
               ("EPS change, 4 weeks", "BEST_EPS_4WK_CHG"), ("EPS change, 3 months", "BEST_EPS_3MO_CHG"), ("EBIT mean", "BEST_EBIT"), ("Free cash flow mean", "BEST_ESTIMATE_FCF")]
    header(ws, 2, ["Annual", "Bloomberg field"] + list(years))
    for i, (label, field) in enumerate(metrics, 3):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=field).font = NOTE
        for j, p in enumerate(years):
            ws.cell(row=i, column=3 + j, value=bdp(GEV, field, p))
    top = len(metrics) + 5
    quarters = ["%dFQ" % k for k in range(1, 7)]
    header(ws, top, ["Quarterly", "Bloomberg field"] + quarters)
    qm = [("Revenue", "BEST_SALES"), ("Revenue number of estimates", "BEST_SALES_NUMEST"), ("EBITDA", "BEST_EBITDA"), ("EBITDA high", "BEST_EBITDA_HI"),
          ("EBITDA low", "BEST_EBITDA_LO"), ("EBITDA number of estimates", "BEST_EBITDA_NUMEST"), ("EPS", "BEST_EPS")]
    for i, (label, field) in enumerate(qm, top + 1):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=field).font = NOTE
        for j, p in enumerate(quarters):
            ws.cell(row=i, column=3 + j, value=bdp(GEV, field, p))
    widths(ws, {"A": 34, "B": 24, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 14})

    start = (2025, 1, 3)
    series = [("GEV EBITDA std dev 2FY (2027)", bdh(GEV, "BEST_EBITDA_STDDEV", start, "2FY")), ("GEV EBITDA high 2FY", bdh(GEV, "BEST_EBITDA_HI", start, "2FY")),
              ("GEV EBITDA low 2FY", bdh(GEV, "BEST_EBITDA_LO", start, "2FY")), ("GEV revenue 1FY (2026)", bdh(GEV, "BEST_SALES", start, "1FY")),
              ("GEV revenue 4FY (2029)", bdh(GEV, "BEST_SALES", start, "4FY")), ("GEV free cash flow 2FY", bdh(GEV, "BEST_ESTIMATE_FCF", start, "2FY"))]
    for name, tic, _ in HYPERSCALERS:
        series.append(("%s capex consensus 3FY" % name, bdh(tic, "BEST_CAPEX", start, "3FY")))
    history_sheet(wb, "EstimateHistory", series, "Weekly history: how wide the 2027 estimates are, and the data-centre builders' capex two years out.")


def bulk(wb):
    ws = wb.create_sheet("Bulk")
    ws.cell(row=1, column=1, value="Each table spills down and to the right from its formula cell.").font = NOTE
    tables = [("GEV suppliers, top 25", GEV, "SUPPLY_CHAIN_SUPPLIERS", "SUPPLY_CHAIN_COUNT_OVERRIDE=25"),
              ("GEV customers, top 25", GEV, "SUPPLY_CHAIN_CUSTOMERS", "SUPPLY_CHAIN_COUNT_OVERRIDE=25"),
              ("Howmet customers, top 25", "HWM US Equity", "SUPPLY_CHAIN_CUSTOMERS", "SUPPLY_CHAIN_COUNT_OVERRIDE=25"),
              ("Howmet past report dates", "HWM US Equity", "EARN_ANN_DT_TIME_HIST_WITH_EPS", None),
              ("Siemens Energy past report dates", "ENR GY Equity", "EARN_ANN_DT_TIME_HIST_WITH_EPS", None)]
    col = 1
    for label, sec, field, ov in tables:
        c = ws.cell(row=2, column=col, value="%s  [%s]" % (label, field))
        c.font, c.fill = HEAD, FILL
        ws.cell(row=3, column=col, value='=BDS("%s","%s"%s)' % (sec, field, ',"%s"' % ov if ov else ""))
        for k in range(8):
            ws.column_dimensions[get_column_letter(col + k)].width = 18
        col += 9


def manual(wb):
    ws = wb.create_sheet("Manual")
    header(ws, 1, ["#", "Type this", "Then", "Save as", "Why"])
    rows = [
        ("1", "HWM US Equity EEO <GO>", "Company Specific tab. Set to quarterly. Actions > Export to Excel.", "31_hwm_company_specific.xlsx",
         "Howmet's gas turbine revenue line: is blade supply growing as fast as GE Vernova's planned output?"),
        ("2", "GEV US Equity EEB <GO>", "Measure EBITDA, period FY2027. Export. Repeat for FY2028.", "32_eeb_2027.xlsx, 33_eeb_2028.xlsx",
         "Each broker's estimate and its date: how fresh the consensus is."),
        ("3", "GEV US Equity MODL <GO>", "Period Q1 2027. Export. Repeat for Q2 2027. If the screen shows how many brokers are behind Power equipment revenue, write it down.",
         "34_modl_2027q1.xlsx, 35_modl_2027q2.xlsx", "The first-half 2027 turbine revenue number and how solid it is."),
        ("4", "GEV US Equity ERN <GO>", "Export the surprise history.", "36_ern.xlsx", "How results compared with estimates, and how the stock moved."),
        ("5", "GEV US Equity SPLC <GO>", "Suppliers view. Export.", "37_splc_suppliers.xlsx", "Quantified supplier links."),
        ("6 (optional)", "GEV US Equity BRC <GO>", "Any recent report whose model shows gigawatts shipped per year or price per kW. Save the PDF.", "38_broker_report.pdf",
         "The Street's own gigawatt assumption."),
        ("7 (optional)", "BNEF <GO>", "Search for a gas plant or turbine order list with turbine maker and start year. Export if found within five minutes.", "39_bnef.xlsx",
         "Deliveries by year from the customer side."),
    ]
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 12, "B": 28, "C": 70, "D": 36, "E": 70})


def main():
    wb = Workbook()
    readme(wb)
    suppliers(wb)
    customers(wb)
    options(wb)
    estimates(wb)
    bulk(wb)
    manual(wb)
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    n = sum(1 for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    print("wrote %s: %d sheets, %d Bloomberg formulas" % (OUT, len(wb.sheetnames), n))


if __name__ == "__main__":
    main()
