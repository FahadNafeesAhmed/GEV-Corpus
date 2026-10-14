"""
Recalculate a workbook in Microsoft Excel (COM) so every formula gets a cached value, then report any error cells.

    python src/excel_recalc.py data/model/GEV_model.xlsx
"""
import os
import sys

import win32com.client
from openpyxl import load_workbook

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")


def recalc(path):
    path = os.path.abspath(path)
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(path)
        xl.CalculateFull()
        wb.Save()
        wb.Close(False)
    finally:
        xl.Quit()


def scrub_names(path):
    """Blank author fields in docProps without re-saving the workbook (keeps Excel's cached values)."""
    import re
    import shutil
    import tempfile
    import zipfile
    tmp = tempfile.mktemp(suffix=".xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                xml = data.decode("utf-8")
                xml = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", "<cp:lastModifiedBy></cp:lastModifiedBy>", xml)
                xml = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator></dc:creator>", xml)
                data = xml.encode("utf-8")
            if item.filename == "xl/workbook.xml":
                xml = data.decode("utf-8")
                xml = re.sub(r"<mc:AlternateContent[^>]*>\s*<mc:Choice Requires=\"x15\">\s*<x15ac:absPath[^>]*/>\s*</mc:Choice>\s*</mc:AlternateContent>", "", xml)
                xml = re.sub(r"<x15ac:absPath[^>]*/>", "", xml)
                data = xml.encode("utf-8")
            if item.filename == "docProps/app.xml":
                xml = data.decode("utf-8")
                xml = re.sub(r"<Company>.*?</Company>", "<Company></Company>", xml)
                xml = re.sub(r"<Manager>.*?</Manager>", "<Manager></Manager>", xml)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


def scan(path):
    vals = load_workbook(path, data_only=True)
    forms = load_workbook(path)
    n_form, errs, empty = 0, [], []
    for ws in forms.worksheets:
        wv = vals[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_form += 1
                    v = wv[c.coordinate].value
                    if isinstance(v, str) and v.strip() in ERRORS:
                        errs.append("%s!%s %s  %s" % (ws.title, c.coordinate, v, c.value))
                    elif v is None:
                        empty.append("%s!%s" % (ws.title, c.coordinate))
    return n_form, errs, empty


if __name__ == "__main__":
    p = sys.argv[1]
    recalc(p)
    scrub_names(p)
    n, errs, empty = scan(p)
    print("formulas: %d | errors: %d | no value: %d" % (n, len(errs), len(empty)))
    for e in errs[:50]:
        print("  ERROR", e)
    for e in empty[:20]:
        print("  EMPTY", e)
