"""
GE Vernova valuation model (12-month horizon), built from filings and the Bloomberg pull of 2026-10-01.

    python src/build_model.py        ->  data/model/GEV_model.xlsx

Conventions: blue = typed input; yellow fill = key assumption the team sets; black = formula; green = link to
another sheet. Base-case drivers default to values that reproduce Bloomberg consensus; change them to your view.
All Street figures are read from the Bloomberg files at build time, not typed by hand.
"""
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
BBG = ROOT / "data" / "bloomberg"
OUT = ROOT / "data" / "model" / "GEV_model.xlsx"

# ---------------------------------------------------------------- styles
F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10, color="000000")
GREEN = Font(name=F, size=10, color="008000")
BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=13, bold=True)
HEAD = Font(name=F, size=10, bold=True, color="FFFFFF")
NOTE = Font(name=F, size=9, italic=True, color="595959")
HEAD_FILL = PatternFill("solid", fgColor="1F3A4D")
SECTION_FILL = PatternFill("solid", fgColor="DCE6EC")
YELLOW = PatternFill("solid", fgColor="FFFF00")
TOP = Border(top=Side(style="thin", color="808080"))

M = '#,##0;(#,##0);"-"'
M1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
X = '0.0"x";(0.0"x");"-"'
PX = '$#,##0.00;($#,##0.00);"-"'
KW = '#,##0;(#,##0);"-"'
GW = '0.0;(0.0);"-"'


def put(ws, ref, value, font=BLACK, fmt=None, fill=None, bold=False, align=None):
    c = ws[ref]
    c.value = value
    c.font = Font(name=F, size=font.size, color=font.color, bold=bold or font.bold, italic=font.italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    return c


def header(ws, row, labels, start_col=1):
    for i, v in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=v)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)


def section(ws, row, text, ncols):
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    c = ws.cell(row=row, column=1, value=text)
    c.font = BOLD


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def title(ws, text, sub):
    put(ws, "A1", text, TITLE)
    put(ws, "A2", sub, NOTE)


# ---------------------------------------------------------------- Bloomberg data (read, not typed)
def load_street():
    q = json.load(open(BBG / "GEV_Pull_2026-10-01" / "segment_quarterly.json", encoding="utf-8"))

    def yr(key, y):
        vals = [q.get("%s|%d Q%d" % (key, y, i)) for i in (1, 2, 3, 4)]
        assert all(v is not None for v in vals), (key, y)
        return sum(vals)

    def h2(key):
        return q["%s|2026 Q3" % key] + q["%s|2026 Q4" % key]

    keys = {
        "p_eq": "Total Revenue|Power|Equipment", "p_sv": "Total Revenue|Power|Services", "p_rev": "Total Revenue||Power",
        "p_e": "Adjusted EBITDA||Power", "e_rev": "Total Revenue||Electrification", "e_e": "Adjusted EBITDA||Electrification",
        "w_rev": "Total Revenue||Wind", "w_e": "Adjusted EBITDA||Wind", "oth": "EBITDA||Other",
        "w_eq": "Total Revenue|Wind|Equipment", "e_eq": "Total Revenue|Electrification|Equipment",
    }
    seg = {"H2 2026": {k: h2(v) for k, v in keys.items()}}
    for y in (2027, 2028, 2029, 2030):
        seg[str(y)] = {k: yr(v, y) for k, v in keys.items()}
    win = {}
    for k in ("p_eq", "w_eq", "e_eq"):
        win[k] = sum(q["%s|%s" % (keys[k], p)] for p in ("2026 Q3", "2026 Q4", "2027 Q1", "2027 Q2"))
    for name, ps in (("p_eq_h2_26", ("2026 Q3", "2026 Q4")), ("p_eq_h1_27", ("2027 Q1", "2027 Q2")), ("p_eq_h2_27", ("2027 Q3", "2027 Q4"))):
        win[name] = sum(q["%s|%s" % (keys["p_eq"], p)] for p in ps)

    wb = load_workbook(BBG / "GEV_bloomberg_pull_filled_2026-10-01.xlsx", data_only=True, read_only=True)
    cs = wb["Consensus"]
    rows = {}
    for r in cs.iter_rows(min_row=2, max_row=18, values_only=True):
        if r[0] and r[1]:
            rows[r[1]] = list(r[2:7])
    annual = {
        "rev": rows["BEST_SALES"], "ebitda": rows["BEST_EBITDA"], "ebitda_med": rows["BEST_EBITDA_MEDIAN"],
        "ebitda_hi": rows["BEST_EBITDA_HI"], "ebitda_lo": rows["BEST_EBITDA_LO"], "ebitda_n": rows["BEST_EBITDA_NUMEST"],
        "ebit": rows["BEST_EBIT"], "eps": rows["BEST_EPS"], "fcf": rows["BEST_ESTIMATE_FCF"],
    }
    snap = wb["Snapshot"]
    peers = []
    gev_snap = None
    for r in snap.iter_rows(min_row=3, max_row=19, values_only=True):
        if not r[0]:
            continue
        rec = {"name": r[0], "ticker": r[1], "px": r[5], "ev": r[7], "fwd_ev_ebitda": r[8], "fwd_pe": r[9],
               "target": r[11], "buy": r[12], "hold": r[13], "sell": r[14], "si_pct": r[16], "shares": r[20]}
        if r[1] == "GEV US Equity":
            gev_snap = rec
        elif r[0] in ("Siemens Energy", "Mitsubishi Heavy", "Eaton", "Vertiv", "Quanta", "Hubbell", "ABB", "Schneider", "Hitachi", "Baker Hughes"):
            peers.append(rec)
    mult = wb["Multiples"]
    import datetime as dt
    pts = []
    for r in mult.iter_rows(min_row=4, max_col=2, values_only=True):
        d, v = r
        if isinstance(d, (dt.datetime, dt.date)) and isinstance(v, (int, float)):
            pts.append((d.date() if isinstance(d, dt.datetime) else d, float(v)))
    last = pts[-1][0]
    yr_pts = [v for d, v in pts if (last - d).days <= 365]
    srt = sorted(yr_pts)
    hist = {"low": min(yr_pts), "high": max(yr_pts), "median": srt[len(srt) // 2] if len(srt) % 2 else (srt[len(srt) // 2 - 1] + srt[len(srt) // 2]) / 2,
            "weeks": len(yr_pts), "weeks_le_26": sum(v <= 26 for v in yr_pts), "last_date": str(last)}
    return seg, win, annual, peers, gev_snap, hist


SEG, WIN, ANN, PEERS, GEVSNAP, MHIST = load_street()

# ---------------------------------------------------------------- historical data (filings)
HIST_COLS = ["FY2023", "FY2024", "FY2025", "H1 2025", "H1 2026"]
H = {
    "eq": [18258, 18952, 20934, 9091, 11713],
    "sv": [14981, 15983, 17134, 8052, 8729],
    "rev_rep": [33239, 34935, 38068, 17143, 20442],
    "p_rev": [17436, 18127, 19767, 9234, 10449],
    "p_eq": [5598, 5708, 6686, 2996, 3851],
    "p_sv": [11838, 12419, 13081, 6238, 6598],
    "e_rev": [6378, 7550, 9642, 4001, 6597],
    "w_rev": [9826, 9701, 9110, 4095, 3459],
    "elim": [-401, -442, -451, -187, -62],
    "p_e": [1722, 2268, 2902, 1303, 1842],
    "e_e": [234, 679, 1433, 519, 1200],
    "w_e": [-1033, -588, -598, -312, -657],
    "corp": [-116, -323, -541, -283, -239],
    "ebitda_rep": [807, 2035, 3196, 1227, 2146],
    "gw": [13.8, 11.9, 15.3, 8.2, 7.5],
}

# ---------------------------------------------------------------- base-case defaults = Bloomberg consensus
FY26 = {k: None for k in ("p_eq", "p_sv", "p_rev", "p_e", "e_rev", "e_e", "w_rev", "w_e", "oth")}
for k in FY26:
    h1 = {"p_eq": 3851, "p_sv": 6598, "p_rev": 10449, "p_e": 1842, "e_rev": 6597, "e_e": 1200, "w_rev": 3459, "w_e": -657, "oth": -239}[k]
    FY26[k] = h1 + SEG["H2 2026"][k]
STREET_Y = {"2026": FY26, "2027": SEG["2027"], "2028": SEG["2028"], "2029": SEG["2029"]}
GW_PATH = {"2026": 17.5, "2027": 21.0, "2028": 24.0, "2029": 26.0}


def defaults():
    d = {}
    prev_sv, prev_e = H["p_sv"][2], H["e_rev"][2]
    for y in ("2026", "2027", "2028", "2029"):
        s = STREET_Y[y]
        d.setdefault("gw", []).append(GW_PATH[y])
        d.setdefault("kw", []).append(round(s["p_eq"] / GW_PATH[y]))
        d.setdefault("svc_g", []).append(round(s["p_sv"] / prev_sv - 1, 3))
        d.setdefault("p_m", []).append(round(s["p_e"] / (s["p_eq"] + s["p_sv"]), 3))
        d.setdefault("e_g", []).append(round(s["e_rev"] / prev_e - 1, 3))
        d.setdefault("e_m", []).append(round(s["e_e"] / s["e_rev"], 3))
        d.setdefault("w_rev", []).append(round(s["w_rev"]))
        d.setdefault("w_e", []).append(round(s["w_e"]))
        d.setdefault("elim", []).append(-400)
        d.setdefault("corp", []).append(round(s["oth"]))
        prev_sv, prev_e = s["p_sv"], s["e_rev"]
    return d


DEF = defaults()

wb = Workbook()
wb.properties.creator = ""
wb.properties.lastModifiedBy = ""
wb.properties.title = "GE Vernova valuation model"

# ======================================================================= History
hs = wb.active
hs.title = "History"
title(hs, "GE Vernova: historical financials ($ millions)", "Blue = typed from filings (source in column G). Black = formula. FY2023-FY2025 on the segment basis before the 1 Jan 2026 realignment; H1 2025 restated and H1 2026 on the new basis.")
header(hs, 4, ["$ millions"] + HIST_COLS + ["Source"])
HR = {}
r = 5
section(hs, r, "Company revenue", 7); r += 1


def hist_row(key, label, src, fmt=M):
    global r
    put(hs, "A%d" % r, label)
    for i, v in enumerate(H[key]):
        put(hs, "%s%d" % ("BCDEF"[i], r), v, BLUE, fmt)
    put(hs, "G%d" % r, src, NOTE)
    HR[key] = r
    r += 1


def hist_formula(key, label, f, fmt=M, bold=False, note=""):
    global r
    put(hs, "A%d" % r, label, BOLD if bold else BLACK)
    for col in "BCDEF":
        put(hs, "%s%d" % (col, r), f(col), BLACK, fmt, bold=bold)
    if note:
        put(hs, "G%d" % r, note, NOTE)
    HR[key] = r
    r += 1


SRC_IS = "Income statement: 10-K FY2024 (0001996810-25-000011) for FY2023; 10-K FY2025 (0001996810-26-000015) p.39; 10-Q Q2 2026 (0001996810-26-000148) p.6"
SRC_SEG = "Segment table: 10-K FY2025 (0001996810-26-000015) MD&A p.26 and Note 24 pp.72-73; 10-Q Q2 2026 (0001996810-26-000148) MD&A p.29 and Note 24 pp.23-24 (H1 2025 restated)"
SRC_PWR = "MD&A Power table (segment revenues incl. intersegment): 10-K FY2025 p.27 for FY2023-FY2025; 10-Q Q2 2026 p.30 for H1 2026 and H1 2025 restated. Note 24 shows external sales only"
hist_row("eq", "Equipment revenue", SRC_IS)
hist_row("sv", "Services revenue", SRC_IS)
hist_formula("rev", "Total revenue", lambda c: "=%s%d+%s%d" % (c, HR["eq"], c, HR["sv"]), bold=True)
hist_row("rev_rep", "Reported total revenue", "Same filings, total revenues line")
hist_formula("chk_rev", "Check (should be 0)", lambda c: "=%s%d-%s%d" % (c, HR["rev"], c, HR["rev_rep"]))
r += 1
section(hs, r, "Segment revenue", 7); r += 1
hist_row("p_rev", "Power revenue", SRC_SEG)
hist_row("p_eq", "   Power equipment", SRC_PWR)
hist_row("p_sv", "   Power services", SRC_PWR)
hist_formula("chk_p", "   Check: equipment + services - Power (rounding +/-1)", lambda c: "=%s%d+%s%d-%s%d" % (c, HR["p_eq"], c, HR["p_sv"], c, HR["p_rev"]))
hist_row("e_rev", "Electrification revenue", SRC_SEG + " (from 2026 includes Prolec GE)")
hist_row("w_rev", "Wind revenue", SRC_SEG)
hist_row("elim", "Eliminations and other", SRC_SEG)
hist_formula("seg_rev", "Segment revenue total", lambda c: "=%s%d+%s%d+%s%d+%s%d" % (c, HR["p_rev"], c, HR["e_rev"], c, HR["w_rev"], c, HR["elim"]), bold=True)
hist_formula("chk_seg", "Check vs reported total (rounding +/-1)", lambda c: "=%s%d-%s%d" % (c, HR["seg_rev"], c, HR["rev_rep"]))
r += 1
section(hs, r, "Segment EBITDA", 7); r += 1
hist_row("p_e", "Power EBITDA", SRC_SEG)
hist_row("e_e", "Electrification EBITDA", SRC_SEG)
hist_row("w_e", "Wind EBITDA", SRC_SEG)
hist_row("corp", "Corporate and other", "Reconciliation of segment EBITDA to net income: 10-K FY2025 p.73 (also MD&A p.26); 10-Q Q2 2026 p.24 (also MD&A p.29)")
hist_formula("ebitda", "Adjusted EBITDA", lambda c: "=SUM(%s%d:%s%d)" % (c, HR["p_e"], c, HR["corp"]), bold=True)
hist_row("ebitda_rep", "Reported adjusted EBITDA", "10-K FY2025 MD&A p.26 (807 / 2,035 / 3,196); 10-Q Q2 2026 MD&A p.29 (1,227 / 2,146). Bloomberg Q2 2026 actual EBITDA (1,250) equals GEV adjusted EBITDA, so consensus is on GEV's definition")
hist_formula("chk_e", "Check (rounding +/-1)", lambda c: "=%s%d-%s%d" % (c, HR["ebitda"], c, HR["ebitda_rep"]))
r += 1
section(hs, r, "Margins and physical output", 7); r += 1
hist_formula("p_m", "Power EBITDA margin", lambda c: "=IF(%s%d=0,0,%s%d/%s%d)" % (c, HR["p_rev"], c, HR["p_e"], c, HR["p_rev"]), PCT)
hist_formula("e_m", "Electrification EBITDA margin", lambda c: "=IF(%s%d=0,0,%s%d/%s%d)" % (c, HR["e_rev"], c, HR["e_e"], c, HR["e_rev"]), PCT)
hist_formula("m", "Adjusted EBITDA margin", lambda c: "=IF(%s%d=0,0,%s%d/%s%d)" % (c, HR["rev"], c, HR["ebitda"], c, HR["rev"]), PCT)
hist_row("gw", "Gas turbine GW shipped", "Sales in units, Gas Turbine Gigawatts: 10-K FY2025 p.26 (FY2023-FY2025); 10-Q Q2 2026 p.29 (H1 2025, H1 2026)", GW)
hist_formula("kw", "Power equipment revenue per kW shipped ($/kW)", lambda c: "=IF(%s%d=0,0,%s%d/%s%d)" % (c, HR["gw"], c, HR["p_eq"], c, HR["gw"]), KW,
             note="All Power equipment over gas GW. Includes non-gas equipment (nuclear, hydro; steam before 2026), estimated about $1.2B a year (range $0.9-1.6B, not disclosed), and aeroderivative units in the GW count")
widths(hs, {"A": 46, "B": 11, "C": 11, "D": 11, "E": 11, "F": 11, "G": 100})
hs.freeze_panes = "B5"

# ======================================================================= Street (Bloomberg)
st = wb.create_sheet("Street")
title(st, "Bloomberg consensus (BEst mean, Standard source), 2026-10-01", "Licensed for academic use. Cite as 'Bloomberg, 2026-10-01'. Annual block = direct annual consensus. Segment block = quarterly consensus summed to years (fewer brokers; does not tie exactly to the annual figures).")
SR = {}
header(st, 4, ["Annual consensus ($ millions)", "", "FY2026", "FY2027", "FY2028", "FY2029", "FY2030", "Bloomberg field"])
r = 5
for key, label, fld, fmt in [("rev", "Revenue", "BEST_SALES", M), ("ebitda", "Adjusted EBITDA (mean)", "BEST_EBITDA", M),
                             ("ebitda_med", "Adjusted EBITDA (median)", "BEST_EBITDA_MEDIAN", M), ("ebitda_hi", "Adjusted EBITDA (high)", "BEST_EBITDA_HI", M),
                             ("ebitda_lo", "Adjusted EBITDA (low)", "BEST_EBITDA_LO", M), ("ebitda_n", "Number of EBITDA estimates", "BEST_EBITDA_NUMEST", M),
                             ("ebit", "EBIT", "BEST_EBIT", M), ("eps", "EPS ($)", "BEST_EPS", '0.00'), ("fcf", "Free cash flow", "BEST_ESTIMATE_FCF", M)]:
    put(st, "A%d" % r, label)
    for i, v in enumerate(ANN[key]):
        put(st, "%s%d" % ("CDEFG"[i], r), round(v, 2) if key == "eps" else round(v), BLUE, fmt)
    put(st, "H%d" % r, fld, NOTE)
    SR["a_" + key] = r
    r += 1
put(st, "A%d" % r, "Adjusted EBITDA margin")
for col in "CDEFG":
    put(st, "%s%d" % (col, r), "=%s%d/%s%d" % (col, SR["a_ebitda"], col, SR["a_rev"]), BLACK, PCT)
SR["a_m"] = r
r += 2

header(st, r, ["Segment consensus ($ millions)", "", "FY2026", "FY2027", "FY2028", "FY2029", "FY2030", "Note"])
r += 1
seg_rows = [("p_eq", "Power equipment revenue"), ("p_sv", "Power services revenue"), ("p_rev", "Power revenue"),
            ("p_e", "Power adjusted EBITDA"), ("e_rev", "Electrification revenue"), ("e_e", "Electrification adjusted EBITDA"),
            ("w_rev", "Wind revenue"), ("w_e", "Wind adjusted EBITDA"), ("oth", "Corporate and other")]
for key, label in seg_rows:
    put(st, "A%d" % r, label)
    vals = [FY26[key]] + [SEG[str(y)][key] for y in (2027, 2028, 2029, 2030)]
    for i, v in enumerate(vals):
        put(st, "%s%d" % ("CDEFG"[i], r), round(v), BLUE, M)
    put(st, "H%d" % r, "FY2026 = H1 2026 actual + H2 2026 consensus" if key != "oth" else "Bloomberg EBITDA 'Other' line; FY2026 = H1 2026 actual (239) + H2 2026 consensus", NOTE)
    SR["s_" + key] = r
    r += 1
put(st, "A%d" % r, "Power margin (on equipment + services)")
for col in "CDEFG":
    put(st, "%s%d" % (col, r), "=IF((%s%d+%s%d)=0,0,%s%d/(%s%d+%s%d))" % (col, SR["s_p_eq"], col, SR["s_p_sv"], col, SR["s_p_e"], col, SR["s_p_eq"], col, SR["s_p_sv"]), BLACK, PCT)
put(st, "H%d" % r, "Bloomberg's Power total comes from a different set of brokers than its equipment/services split, so they do not add up exactly; the model uses equipment + services", NOTE)
SR["s_p_m"] = r
r += 1
for key, label, num, den in [("s_e_m", "Electrification margin", "s_e_e", "s_e_rev"), ("s_w_m", "Wind margin", "s_w_e", "s_w_rev")]:
    put(st, "A%d" % r, label)
    for col in "CDEFG":
        put(st, "%s%d" % (col, r), "=IF(%s%d=0,0,%s%d/%s%d)" % (col, SR[den], col, SR[num], col, SR[den]), BLACK, PCT)
    SR[key] = r
    r += 1
put(st, "A%d" % r, "Sum of segment EBITDA", BOLD)
for col in "CDEFG":
    put(st, "%s%d" % (col, r), "=%s%d+%s%d+%s%d+%s%d" % (col, SR["s_p_e"], col, SR["s_e_e"], col, SR["s_w_e"], col, SR["s_oth"]), BLACK, M, bold=True)
SR["s_sum"] = r
r += 2

header(st, r, ["Valuation and positioning", "Value", "", "", "", "", "", "Note"])
r += 1
for key, label, v, fmt, note in [
    ("fwd", "GEV forward EV/EBITDA, next 12 months", GEVSNAP["fwd_ev_ebitda"], X, "Bloomberg BEST_EV_TO_BEST_EBITDA, blended forward"),
    ("low", "Lowest weekly value, last 12 months", MHIST["low"], X, "%d weeks to %s; weeks at or below 26x: %d" % (MHIST["weeks"], MHIST["last_date"], MHIST["weeks_le_26"])),
    ("med", "Median weekly value, last 12 months", MHIST["median"], X, ""),
    ("high", "Highest weekly value, last 12 months", MHIST["high"], X, ""),
    ("fpe", "GEV forward P/E, next 12 months", GEVSNAP["fwd_pe"], X, ""),
    ("target", "Average analyst price target ($)", GEVSNAP["target"], PX, "Buy %d / Hold %d / Sell %d" % (GEVSNAP["buy"], GEVSNAP["hold"], GEVSNAP["sell"])),
    ("si", "Short interest, % of float", GEVSNAP["si_pct"] / 100.0, PCT, ""),
]:
    put(st, "A%d" % r, label)
    put(st, "B%d" % r, round(v, 4) if isinstance(v, float) else v, BLUE, fmt)
    put(st, "H%d" % r, note, NOTE)
    SR["v_" + key] = r
    r += 1
r += 1
header(st, r, ["Peers: forward EV/EBITDA", "Multiple", "Ticker", "", "", "", "", "Bloomberg BEST_EV_TO_BEST_EBITDA, blended forward, Snapshot 2026-10-01. Bloom Energy (51.7x) excluded as an outlier"])
r += 1
SR["peer_first"] = r
for p in PEERS:
    put(st, "A%d" % r, p["name"])
    put(st, "B%d" % r, round(p["fwd_ev_ebitda"], 2), BLUE, X)
    put(st, "C%d" % r, p["ticker"], NOTE)
    r += 1
SR["peer_last"] = r - 1
put(st, "A%d" % r, "Peer median", BOLD)
put(st, "B%d" % r, "=MEDIAN(B%d:B%d)" % (SR["peer_first"], SR["peer_last"]), BLACK, X, bold=True)
SR["peer_med"] = r
widths(st, {"A": 40, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 70})
st.freeze_panes = "C5"

# ======================================================================= Drivers
dr = wb.create_sheet("Drivers")
title(dr, "Forecast drivers: the assumptions you set", "Yellow cells are your assumptions. Defaults reproduce Bloomberg segment consensus (columns H-K) so the base case starts at the Street; replace them with your own view and note why in column L. Eliminations default to (400), the historical level.")
header(dr, 4, ["Driver", "Unit", "FY2026E", "FY2027E", "FY2028E", "FY2029E", "", "Street FY2026", "Street FY2027", "Street FY2028", "Street FY2029", "Evidence and source"])
DR = {}
r = 5
YC = "CDEF"
SC = "HIJK"
STREET_COL = {"H": "C", "I": "D", "J": "E", "K": "F"}   # Drivers street column -> Street sheet column (FY2026..FY2029)


def driver(key, label, unit, fmt, street_f, evidence):
    global r
    put(dr, "A%d" % r, label)
    put(dr, "B%d" % r, unit, NOTE)
    for i, col in enumerate(YC):
        put(dr, "%s%d" % (col, r), DEF[key][i], BLUE, fmt, fill=YELLOW)
    for col in SC:
        f = street_f(STREET_COL[col], col)
        if f:
            put(dr, "%s%d" % (col, r), f, GREEN, fmt)
    put(dr, "L%d" % r, evidence, NOTE)
    DR[key] = r
    r += 1


section(dr, r, "Power", 12); r += 1
driver("gw", "Gas turbine GW shipped", "GW", GW, lambda s, c: None,
       "Company: 5 GW a quarter from Q3 2026 (H1 2026 actual 7.5); 6 GW a quarter from sometime in H2 2027, held for a period into 2028, then a step to 7 GW a quarter (CEO, 2026-09-16, no dates given); 24 GW of annual output in 2028 (Q2 2026 release; an output rate, not a stated delivery total); actions toward 30 GW in 2030. Highest year so far 15.3 GW (2025)")
driver("kw", "Power equipment revenue per kW shipped", "$/kW", KW, lambda s, c: "=IF(%s%d=0,0,Street!%s%d/%s%d)" % ({"H": "C", "I": "D", "J": "E", "K": "F"}[c], DR["gw"], s, SR["s_p_eq"], {"H": "C", "I": "D", "J": "E", "K": "F"}[c], DR["gw"]),
       "Realised: 406 (2023), 480 (2024), 437 (2025), 513 (H1 2026). Signed Power backlog per gas GW 741 at 6/30/26 ($39,261M / 53 GW, includes non-gas). Power equipment orders per gas GW ordered 753 (Q4 2025), 705 (Q1 2026), 1,049 (Q2 2026, lifted by aeroderivative mix). All-in figures, not like-for-like turbine prices. Street value shown = Street revenue / your GW")
driver("svc_g", "Power services revenue growth", "%", PCT, lambda s, c: "=Street!%s%d/%s-1" % (s, SR["s_p_sv"], "History!D%d" % HR["p_sv"] if c == "H" else "Street!%s%d" % (chr(ord(s) - 1), SR["s_p_sv"])),
       "Power services grew +4.9% (2024), +5.3% (2025); H1 2026 +5.8% on the restated basis. Services RPO $72.4B for Power. FY2026 growth is measured on the old-basis FY2025 figure")
driver("p_m", "Power EBITDA margin", "%", PCT, lambda s, c: "=Street!%s%d" % (s, SR["s_p_m"]),
       "Actual 9.9% (2023), 12.5% (2024), 14.7% (2025), 17.6% (H1 2026). FY2026 guide 17-19%. 2028 target 22%. Street margin computed on equipment + services, the base the forecast uses")
section(dr, r, "Electrification", 12); r += 1
driver("e_g", "Electrification revenue growth", "%", PCT, lambda s, c: "=Street!%s%d/%s-1" % (s, SR["s_e_rev"], "History!D%d" % HR["e_rev"] if c == "H" else "Street!%s%d" % (chr(ord(s) - 1), SR["s_e_rev"])),
       "FY2026 includes about $3.1B from Prolec GE (acquired 2026-02-02); FY2026 guide $14.5-15.0B. Equipment RPO $40.6B at 6/30/26. FY2026 growth is measured on the old-basis FY2025 figure (the realignment moved some revenue out of Electrification)")
driver("e_m", "Electrification EBITDA margin", "%", PCT, lambda s, c: "=Street!%s%d" % (s, SR["s_e_m"]),
       "Actual 3.7% (2023), 9.0% (2024), 14.9% (2025), 18.2% (H1 2026). FY2026 guide 18-20%. 2028 target 22%")
section(dr, r, "Wind, eliminations, corporate", 12); r += 1
driver("w_rev", "Wind revenue", "$M", M, lambda s, c: "=Street!%s%d" % (s, SR["s_w_rev"]),
       "Actual 9,701 (2024), 9,110 (2025), 3,459 (H1 2026). FY2026 guide: organic down low double digits")
driver("w_e", "Wind EBITDA", "$M", M, lambda s, c: "=Street!%s%d" % (s, SR["s_w_e"]),
       "Actual (588) 2024, (598) 2025, (657) H1 2026. FY2026 guide about (400), so H2 needs about +257. Best quarter since 2023: +19. 2028 target 6% margin")
driver("elim", "Eliminations and other revenue", "$M", M, lambda s, c: None,
       "History (401), (442), (451); H1 2026 (62)")
driver("corp", "Corporate and other EBITDA", "$M", M, lambda s, c: "=Street!%s%d" % (s, SR["s_oth"]),
       "Actual (116), (323), (541), H1 2026 (239). FY2026 guide: cost of $450-500M (midpoint 475); default uses the Street figure (490)")
r += 1
section(dr, r, "Company guidance and targets (used for checks and the bear case)", 12); r += 1
GUIDE = [("g_rev_lo", "FY2026 revenue guide, low", 45500, M), ("g_rev_hi", "FY2026 revenue guide, high", 46500, M),
         ("g_m_lo", "FY2026 adjusted EBITDA margin guide, low", 0.12, PCT), ("g_m_hi", "FY2026 adjusted EBITDA margin guide, high", 0.14, PCT),
         ("f_rev", "2028 framework: revenue", 56000, M), ("f_m", "2028 framework: adjusted EBITDA margin", 0.20, PCT),
         ("f_pm", "2028 target: Power margin", 0.22, PCT), ("f_em", "2028 target: Electrification margin", 0.22, PCT), ("f_wm", "2028 target: Wind margin", 0.06, PCT)]
for key, label, v, fmt in GUIDE:
    put(dr, "A%d" % r, label)
    put(dr, "C%d" % r, v, BLUE, fmt)
    put(dr, "L%d" % r, "Q2 2026 release (0001996810-26-000147), guidance issued 2026-07-22" if key.startswith("g_") else "Q4 2025 release (0001996810-26-000012), framework issued 2026-01-28", NOTE)
    DR[key] = r
    r += 1
widths(dr, {"A": 40, "B": 7, "C": 11, "D": 11, "E": 11, "F": 11, "G": 2, "H": 12, "I": 12, "J": 12, "K": 12, "L": 120})
dr.freeze_panes = "C5"

# ======================================================================= Forecast
fc = wb.create_sheet("Forecast")
title(fc, "Forecast by segment ($ millions)", "Every number is a formula on Drivers and History. Change the yellow cells on Drivers, not this sheet.")
header(fc, 4, ["$ millions", "FY2025A", "FY2026E", "FY2027E", "FY2028E", "FY2029E", "Note"])
FR = {}
r = 5
YEARS = "CDEF"


def frow(key, label, hist_ref, fcast, fmt=M, bold=False, note="", font_hist=GREEN):
    global r
    put(fc, "A%d" % r, label, BOLD if bold else BLACK)
    if hist_ref is not None:
        put(fc, "B%d" % r, hist_ref, font_hist, fmt, bold=bold)
    for i, col in enumerate(YEARS):
        prev = "B" if i == 0 else YEARS[i - 1]
        f = fcast(col, prev)
        is_link = f.startswith("=Drivers!") and "*" not in f and "+" not in f
        put(fc, "%s%d" % (col, r), f, GREEN if is_link else BLACK, fmt, bold=bold)
    if note:
        put(fc, "G%d" % r, note, NOTE)
    FR[key] = r
    r += 1


section(fc, r, "Power", 7); r += 1
frow("gw", "Gas turbine GW shipped", "=History!D%d" % HR["gw"], lambda c, p: "=Drivers!%s%d" % (c, DR["gw"]), GW)
frow("kw", "Revenue per kW shipped ($/kW)", "=History!D%d" % HR["kw"], lambda c, p: "=Drivers!%s%d" % (c, DR["kw"]), KW)
frow("p_eq", "Power equipment revenue", "=History!D%d" % HR["p_eq"], lambda c, p: "=%s%d*%s%d" % (c, FR["gw"], c, FR["kw"]), note="GW x $/kW (1 GW = 1,000,000 kW, so GW x $/kW = $ millions)")
frow("p_sv", "Power services revenue", "=History!D%d" % HR["p_sv"], lambda c, p: "=%s%d*(1+Drivers!%s%d)" % (p, r, c, DR["svc_g"]))
frow("p_rev", "Power revenue", "=History!D%d" % HR["p_rev"], lambda c, p: "=%s%d+%s%d" % (c, FR["p_eq"], c, FR["p_sv"]), bold=True)
frow("p_m", "Power EBITDA margin", "=History!D%d" % HR["p_m"], lambda c, p: "=Drivers!%s%d" % (c, DR["p_m"]), PCT)
frow("p_e", "Power EBITDA", "=History!D%d" % HR["p_e"], lambda c, p: "=%s%d*%s%d" % (c, FR["p_rev"], c, FR["p_m"]), bold=True)
r += 1
section(fc, r, "Electrification", 7); r += 1
frow("e_rev", "Electrification revenue", "=History!D%d" % HR["e_rev"], lambda c, p: "=%s%d*(1+Drivers!%s%d)" % (p, r, c, DR["e_g"]), bold=True)
frow("e_m", "Electrification EBITDA margin", "=History!D%d" % HR["e_m"], lambda c, p: "=Drivers!%s%d" % (c, DR["e_m"]), PCT)
frow("e_e", "Electrification EBITDA", "=History!D%d" % HR["e_e"], lambda c, p: "=%s%d*%s%d" % (c, FR["e_rev"], c, FR["e_m"]), bold=True)
r += 1
section(fc, r, "Wind, eliminations, corporate", 7); r += 1
frow("w_rev", "Wind revenue", "=History!D%d" % HR["w_rev"], lambda c, p: "=Drivers!%s%d" % (c, DR["w_rev"]))
frow("w_e", "Wind EBITDA", "=History!D%d" % HR["w_e"], lambda c, p: "=Drivers!%s%d" % (c, DR["w_e"]))
frow("elim", "Eliminations and other revenue", "=History!D%d" % HR["elim"], lambda c, p: "=Drivers!%s%d" % (c, DR["elim"]))
frow("corp", "Corporate and other EBITDA", "=History!D%d" % HR["corp"], lambda c, p: "=Drivers!%s%d" % (c, DR["corp"]))
r += 1
section(fc, r, "Company", 7); r += 1
frow("rev", "Total revenue", "=History!D%d" % HR["rev"], lambda c, p: "=%s%d+%s%d+%s%d+%s%d" % (c, FR["p_rev"], c, FR["e_rev"], c, FR["w_rev"], c, FR["elim"]), bold=True)
put(fc, "A%d" % r, "Revenue growth")
for i, col in enumerate("BCDEF"):
    if col == "B":
        put(fc, "B%d" % r, "=History!D%d/History!C%d-1" % (HR["rev"], HR["rev"]), GREEN, PCT)
    else:
        put(fc, "%s%d" % (col, r), "=%s%d/%s%d-1" % (col, FR["rev"], "BCDEF"[i - 1], FR["rev"]), BLACK, PCT)
FR["rev_g"] = r
r += 1
frow("ebitda", "Adjusted EBITDA", "=History!D%d" % HR["ebitda"], lambda c, p: "=%s%d+%s%d+%s%d+%s%d" % (c, FR["p_e"], c, FR["e_e"], c, FR["w_e"], c, FR["corp"]), bold=True)
frow("m", "Adjusted EBITDA margin", "=History!D%d" % HR["m"], lambda c, p: "=IF(%s%d=0,0,%s%d/%s%d)" % (c, FR["rev"], c, FR["ebitda"], c, FR["rev"]), PCT, font_hist=GREEN)
put(fc, "A%d" % r, "Adjusted EBITDA growth")
put(fc, "B%d" % r, "=History!D%d/History!C%d-1" % (HR["ebitda"], HR["ebitda"]), GREEN, PCT)
for i, col in enumerate(YEARS):
    put(fc, "%s%d" % (col, r), "=%s%d/%s%d-1" % (col, FR["ebitda"], ("B" if i == 0 else YEARS[i - 1]), FR["ebitda"]), BLACK, PCT)
FR["e_g"] = r
r += 2
section(fc, r, "Against the Street (Bloomberg, 2026-10-01)", 7); r += 1
put(fc, "A%d" % r, "Street adjusted EBITDA (annual mean)")
put(fc, "B%d" % r, "=History!D%d" % HR["ebitda_rep"], GREEN, M)
for col, sc in zip(YEARS, "CDEF"):
    put(fc, "%s%d" % (col, r), "=Street!%s%d" % (sc, SR["a_ebitda"]), GREEN, M)
FR["st_e"] = r
r += 1
put(fc, "A%d" % r, "Model vs Street")
for col in YEARS:
    put(fc, "%s%d" % (col, r), "=%s%d/%s%d-1" % (col, FR["ebitda"], col, FR["st_e"]), BLACK, PCT)
put(fc, "G%d" % r, "With default drivers the difference is between Bloomberg segment consensus and annual consensus (different broker sets), not a view", NOTE)
FR["vs_e"] = r
r += 1
put(fc, "A%d" % r, "Street revenue (annual mean)")
for col, sc in zip(YEARS, "CDEF"):
    put(fc, "%s%d" % (col, r), "=Street!%s%d" % (sc, SR["a_rev"]), GREEN, M)
FR["st_r"] = r
r += 1
put(fc, "A%d" % r, "Model vs Street")
for col in YEARS:
    put(fc, "%s%d" % (col, r), "=%s%d/%s%d-1" % (col, FR["rev"], col, FR["st_r"]), BLACK, PCT)
FR["vs_r"] = r
r += 2
section(fc, r, "Checks against company guidance", 7); r += 1
put(fc, "A%d" % r, "FY2026 revenue inside the $45.5-46.5B guide")
put(fc, "C%d" % r, '=IF(AND(C%d>=Drivers!C%d,C%d<=Drivers!C%d),"Inside guide","Outside guide")' % (FR["rev"], DR["g_rev_lo"], FR["rev"], DR["g_rev_hi"]))
FR["chk_g1"] = r
r += 1
put(fc, "A%d" % r, "FY2026 margin inside the 12-14% guide")
put(fc, "C%d" % r, '=IF(AND(C%d>=Drivers!C%d,C%d<=Drivers!C%d),"Inside guide","Outside guide")' % (FR["m"], DR["g_m_lo"], FR["m"], DR["g_m_hi"]))
FR["chk_g2"] = r
r += 1
put(fc, "A%d" % r, "FY2028 EBITDA vs 2028 framework ($56B x 20%)")
put(fc, "E%d" % r, "=E%d/(Drivers!C%d*Drivers!C%d)-1" % (FR["ebitda"], DR["f_rev"], DR["f_m"]), BLACK, PCT)
FR["chk_f"] = r
r += 1
widths(fc, {"A": 40, "B": 11, "C": 11, "D": 11, "E": 11, "F": 11, "G": 70})
fc.freeze_panes = "B5"

# ======================================================================= Valuation
va = wb.create_sheet("Valuation")
title(va, "Valuation: price in about 12 months on forward EV/EBITDA", "Price = (exit multiple x next-twelve-month EBITDA + net cash bridge) / shares. Yellow cells are your judgments.")
VR = {}
r = 4
header(va, r, ["Market data and bridge", "Value", "", "", "", "", "", "Source"])
r += 1


def vin(key, label, v, fmt, src, fill=None, font=BLUE):
    global r
    put(va, "A%d" % r, label)
    put(va, "B%d" % r, v, font, fmt, fill=fill)
    put(va, "H%d" % r, src, NOTE)
    VR[key] = r
    r += 1


vin("px", "Share price ($)", round(GEVSNAP["px"], 2), PX, "Bloomberg, 2026-10-01 (last price in the pull)")
vin("sh", "Shares outstanding, basic (millions)", 266.334, '#,##0.000', "10-Q Q2 2026 cover: 266,333,581 shares; Bloomberg 266.334. Used for today's market value")
vin("dil", "Dilutive shares (millions)", 3, '#,##0.000', "10-Q Q2 2026 Note 18: dilutive effect of common stock equivalents, 3 million")
put(va, "A%d" % r, "Diluted shares for target prices (millions)")
put(va, "B%d" % r, "=B%d+B%d" % (VR["sh"], VR["dil"]), BLACK, '#,##0.000')
VR["dsh"] = r
r += 1
vin("cash", "Cash, cash equivalents and restricted cash", 13120, M, "10-Q Q2 2026, cash flow statement, 6/30/26")
vin("restr", "Less restricted cash", -400, M, "10-Q Q2 2026 MD&A, Capital Resources and Liquidity: '$13.1 billion, $0.4 billion of which was restricted use cash'")
vin("debt", "Total borrowings", -2849, M, "10-Q Q2 2026 Note 14 (long-term 2,794 + current 55)")
vin("nci", "Noncontrolling interests", -1158, M, "10-Q Q2 2026, statement of equity")
vin("build", "Cash added between 6/30/26 and the target date", 0, M, "Your assumption. 0 = static bridge (conservative). H1 2026 FCF is already in the 6/30/26 cash. A full estimate: H2 2026 FCF (guide midpoint 12.0B less H1 actual 9.9B = about 2.1B) + 0.75 x Street FY2027 FCF 7.65B (about 5.7B) less dividends (about 0.7B) = about 7.2B (about $27 a share). If you subtract buybacks, also cut the share count", fill=YELLOW)
put(va, "A%d" % r, "Net cash bridge", BOLD)
put(va, "B%d" % r, "=B%d+B%d+B%d+B%d" % (VR["cash"], VR["restr"], VR["debt"], VR["nci"]), BLACK, M, bold=True)
VR["bridge"] = r
r += 1
put(va, "A%d" % r, "Market capitalisation")
put(va, "B%d" % r, "=B%d*B%d" % (VR["px"], VR["sh"]), BLACK, M)
VR["mcap"] = r
r += 1
put(va, "A%d" % r, "Enterprise value today")
put(va, "B%d" % r, "=B%d-B%d" % (VR["mcap"], VR["bridge"]), BLACK, M)
VR["ev"] = r
r += 1
put(va, "A%d" % r, "Next-twelve-month EBITDA today (25% FY2026 + 75% FY2027, model)")
put(va, "B%d" % r, "=0.25*Forecast!C%d+0.75*Forecast!D%d" % (FR["ebitda"], FR["ebitda"]), GREEN, M)
VR["ntm_now"] = r
r += 1
put(va, "A%d" % r, "Forward EV/EBITDA today (model basis)")
put(va, "B%d" % r, "=B%d/B%d" % (VR["ev"], VR["ntm_now"]), BLACK, X)
VR["m_now"] = r
r += 1
put(va, "A%d" % r, "Forward EV/EBITDA today (Bloomberg basis)")
put(va, "B%d" % r, "=Street!B%d" % SR["v_fwd"], GREEN, X)
put(va, "H%d" % r, "EV agrees with Bloomberg within 0.3% (Bloomberg 254,694). The gap is in next-twelve-month EBITDA: Bloomberg about 8,923 (quarterly, Q4-heavy) vs the model's calendar 25/75 blend. The return split below uses the Bloomberg basis", NOTE)
VR["m_bbg"] = r
r += 2

header(va, r, ["Target date", "Value", "", "", "", "", "", "Note"])
r += 1
vin("months", "Months of FY2028 inside the next twelve months at the target date", 9, '0', "9 = target date 30 Sep 2027 (12 months out). 0 = 31 Dec 2026; 1 = late January 2027. Calendar weights; GEV earns more in Q4, so this overstates target-date EBITDA by about 0.8% (about $9 a share)", fill=YELLOW)
put(va, "A%d" % r, "Weight on FY2028")
put(va, "B%d" % r, "=B%d/12" % VR["months"], BLACK, PCT)
VR["w"] = r
r += 2

header(va, r, ["Scenarios", "Bear", "Base", "Bull", "", "", "", "How each is built"])
r += 1
put(va, "A%d" % r, "FY2027 adjusted EBITDA")
put(va, "B%d" % r, "=((Drivers!C%d+Drivers!C%d)/2*(Drivers!C%d+Drivers!C%d)/2+Drivers!C%d*Drivers!C%d)/2" % (DR["g_rev_lo"], DR["g_rev_hi"], DR["g_m_lo"], DR["g_m_hi"], DR["f_rev"], DR["f_m"]), BLACK, M)
put(va, "C%d" % r, "=Forecast!D%d" % FR["ebitda"], GREEN, M)
VR["s27"] = r
r += 1
put(va, "A%d" % r, "FY2028 adjusted EBITDA")
put(va, "B%d" % r, "=Drivers!C%d*Drivers!C%d" % (DR["f_rev"], DR["f_m"]), BLACK, M)
put(va, "C%d" % r, "=Forecast!E%d" % FR["ebitda"], GREEN, M)
VR["s28"] = r
r += 1
vin_row = r
put(va, "A%d" % r, "Bull: EBITDA above base")
put(va, "D%d" % r, 0.05, BLUE, PCT, fill=YELLOW)
put(va, "H%d" % r, "Your judgment. Context: GEV beat pre-report revenue consensus by about 5% over the last four reports (+8.7%, +7.9%, +0.6%, +3.3%); margin beats were smaller. The backlog schedule sits 4.5% above Street equipment revenue for Q3 2026-Q2 2027, worth only about $0.2B of EBITDA (Backlog B17) and in an earlier window", NOTE)
VR["uplift"] = r
r += 1
put(va, "D%d" % VR["s27"], "=C%d*(1+D%d)" % (VR["s27"], VR["uplift"]), BLACK, M)
put(va, "D%d" % VR["s28"], "=C%d*(1+D%d)" % (VR["s28"], VR["uplift"]), BLACK, M)
put(va, "H%d" % VR["s27"], "Bear = GEV's own path: midpoint of FY2026 guide and 2028 framework (an interpolation, not guidance). Base = your forecast on Drivers (defaults reproduce consensus). Bull = base x (1 + uplift)", NOTE)
put(va, "H%d" % VR["s28"], "Bear = 2028 framework ($56B x 20%). Base = your forecast", NOTE)
put(va, "A%d" % r, "Next-twelve-month EBITDA at the target date", BOLD)
for col in "BCD":
    put(va, "%s%d" % (col, r), "=(1-$B$%d)*%s%d+$B$%d*%s%d" % (VR["w"], col, VR["s27"], VR["w"], col, VR["s28"]), BLACK, M, bold=True)
VR["ntm"] = r
r += 1
put(va, "A%d" % r, "Exit multiple (EV / next-twelve-month EBITDA)")
for col, v in zip("BCD", (21.0, 25.0, 28.5)):
    put(va, "%s%d" % (col, r), v, BLUE, X, fill=YELLOW)
put(va, "H%d" % r, "Evidence: today 28.5x (Bloomberg); 12-month low 27.3x; 2025 sell-off about 18-24x; Eaton 21.1x, Vertiv 20.6x, Quanta 22.3x, Hubbell 15.6x; peer median on Street sheet. EBITDA growth slows from about 40% a year to about 25-30% by late 2027", NOTE)
VR["mult"] = r
r += 1
put(va, "A%d" % r, "Enterprise value")
for col in "BCD":
    put(va, "%s%d" % (col, r), "=%s%d*%s%d" % (col, VR["mult"], col, VR["ntm"]), BLACK, M)
VR["sev"] = r
r += 1
put(va, "A%d" % r, "Plus net cash bridge and cash added")
for col in "BCD":
    put(va, "%s%d" % (col, r), "=$B$%d+$B$%d" % (VR["bridge"], VR["build"]), BLACK, M)
VR["sbr"] = r
r += 1
put(va, "A%d" % r, "Equity value")
for col in "BCD":
    put(va, "%s%d" % (col, r), "=%s%d+%s%d" % (col, VR["sev"], col, VR["sbr"]), BLACK, M)
VR["seq"] = r
r += 1
put(va, "A%d" % r, "Price per share ($)", BOLD)
for col in "BCD":
    put(va, "%s%d" % (col, r), "=%s%d/$B$%d" % (col, VR["seq"], VR["dsh"]), BLACK, PX, bold=True)
VR["sp"] = r
r += 1
put(va, "A%d" % r, "Return from today's price", BOLD)
for col in "BCD":
    put(va, "%s%d" % (col, r), "=%s%d/$B$%d-1" % (col, VR["sp"], VR["px"]), BLACK, PCT, bold=True)
VR["sr"] = r
r += 1
put(va, "A%d" % r, "Probability")
for col, v in zip("BCD", (0.25, 0.50, 0.25)):
    put(va, "%s%d" % (col, r), v, BLUE, PCT, fill=YELLOW)
put(va, "H%d" % r, "Your judgment; must sum to 100%", NOTE)
VR["prob"] = r
r += 1
put(va, "A%d" % r, "Probability check (should be 100%)")
put(va, "B%d" % r, "=SUM(B%d:D%d)" % (VR["prob"], VR["prob"]), BLACK, PCT)
VR["pchk"] = r
r += 1
put(va, "A%d" % r, "Probability-weighted price ($)", BOLD)
put(va, "B%d" % r, "=SUMPRODUCT(B%d:D%d,B%d:D%d)" % (VR["sp"], VR["sp"], VR["prob"], VR["prob"]), BLACK, PX, bold=True)
VR["pw"] = r
r += 1
put(va, "A%d" % r, "Probability-weighted return", BOLD)
put(va, "B%d" % r, "=B%d/B%d-1" % (VR["pw"], VR["px"]), BLACK, PCT, bold=True)
VR["pwr"] = r
r += 2

header(va, r, ["What the base case needs", "Value", "", "", "", "", "", "Note"])
r += 1
put(va, "A%d" % r, "Break-even exit multiple (base): price stays at today's level")
put(va, "B%d" % r, "=($B$%d*$B$%d-$B$%d-$B$%d)/C%d" % (VR["px"], VR["dsh"], VR["bridge"], VR["build"], VR["ntm"]), BLACK, X)
put(va, "H%d" % r, "Below this multiple at the target date the long loses money on base-case EBITDA", NOTE)
VR["be"] = r
r += 1
put(va, "A%d" % r, "Price if today's multiple (Bloomberg basis) is kept")
put(va, "B%d" % r, "=($B$%d*C%d+$B$%d+$B$%d)/$B$%d" % (VR["m_bbg"], VR["ntm"], VR["bridge"], VR["build"], VR["dsh"]), BLACK, PX)
VR["p_keep"] = r
r += 1
put(va, "A%d" % r, "Return from EBITDA growth and cash added (multiple unchanged)")
put(va, "B%d" % r, "=B%d/$B$%d-1" % (VR["p_keep"], VR["px"]), BLACK, PCT)
VR["r_growth"] = r
r += 1
put(va, "A%d" % r, "Return from the change in multiple")
put(va, "B%d" % r, "=C%d-B%d" % (VR["sr"], VR["r_growth"]), BLACK, PCT)
put(va, "H%d" % r, "Base return = growth part + multiple part", NOTE)
VR["r_mult"] = r
r += 1
put(va, "A%d" % r, "Base EBITDA at target date vs Street")
put(va, "B%d" % r, "=C%d/((1-$B$%d)*Street!D%d+$B$%d*Street!E%d)-1" % (VR["ntm"], VR["w"], SR["a_ebitda"], VR["w"], SR["a_ebitda"]), BLACK, PCT)
put(va, "H%d" % r, "Size of your estimate difference against Bloomberg consensus", NOTE)
VR["vs_st"] = r
r += 2

header(va, r, ["Sensitivity: price per share", "EBITDA -10%", "-5%", "Base", "+5%", "+10%"])
r += 1
put(va, "A%d" % r, "Change in base next-twelve-month EBITDA", NOTE)
for col, v in zip("BCDEF", (-0.10, -0.05, 0.0, 0.05, 0.10)):
    put(va, "%s%d" % (col, r), v, BLUE, PCT)
VR["sens_h"] = r
r += 1
VR["sens_first"] = r
for m in (19, 21, 23, 25, 27, 29):
    put(va, "A%d" % r, m, BLUE, X)
    for col in "BCDEF":
        put(va, "%s%d" % (col, r), "=($A%d*$C$%d*(1+%s$%d)+$B$%d+$B$%d)/$B$%d" % (r, VR["ntm"], col, VR["sens_h"], VR["bridge"], VR["build"], VR["dsh"]), BLACK, PX)
    r += 1
VR["sens_last"] = r - 1
put(va, "A%d" % r, "Rows: exit multiple. Columns: base next-twelve-month EBITDA changed by the percentage shown.", NOTE)
widths(va, {"A": 58, "B": 13, "C": 13, "D": 13, "E": 13, "F": 13, "G": 2, "H": 110})
va.freeze_panes = "B4"

# ======================================================================= Backlog
bk = wb.create_sheet("Backlog")
title(bk, "Backlog checks: what is already signed, against the Street", "Arithmetic on disclosed backlog. Supporting evidence for the forecast, not a forecast itself.")
BR = {}
r = 4
header(bk, r, ["A. Equipment revenue scheduled from backlog, next four quarters", "Value", "", "", "", "", "", "Source"])
r += 1


def bin_(key, label, v, fmt, src, fill=None):
    global r
    put(bk, "A%d" % r, label)
    put(bk, "B%d" % r, v, BLUE, fmt, fill=fill)
    put(bk, "H%d" % r, src, NOTE)
    BR[key] = r
    r += 1


bin_("rpo", "Equipment backlog (RPO), 6/30/26", 87821, M, "10-Q Q2 2026 (0001996810-26-000148), remaining performance obligations note")
bin_("share", "Share expected within one year", 0.36, PCT, "Same note: 36% / 65% / 97% within 1 / 2 / 5 years")
put(bk, "A%d" % r, "Scheduled equipment revenue, Q3 2026 to Q2 2027", BOLD)
put(bk, "B%d" % r, "=B%d*B%d" % (BR["rpo"], BR["share"]), BLACK, M, bold=True)
BR["sched"] = r
r += 1
bin_("st_p", "Street: Power equipment, Q3 2026 to Q2 2027", round(WIN["p_eq"]), M, "Bloomberg, 2026-10-01: quarterly division consensus")
bin_("st_w", "Street: Wind equipment, same quarters", round(WIN["w_eq"]), M, "Bloomberg, 2026-10-01")
bin_("st_e", "Street: Electrification equipment, same quarters", round(WIN["e_eq"]), M, "Bloomberg, 2026-10-01. Division figures are before intersegment eliminations (about 0.5%), so the gap below is slightly understated")
put(bk, "A%d" % r, "Street equipment revenue, Q3 2026 to Q2 2027", BOLD)
put(bk, "B%d" % r, "=SUM(B%d:B%d)" % (BR["st_p"], BR["st_e"]), BLACK, M, bold=True)
BR["st"] = r
r += 1
put(bk, "A%d" % r, "Schedule above the Street (no new orders counted)")
put(bk, "B%d" % r, "=B%d/B%d-1" % (BR["sched"], BR["st"]), BLACK, PCT)
BR["gap"] = r
r += 1
bin_("uplift", "Ratio of actual to scheduled (your assumption)", 1.00, '0.000', "Seven reported windows: 0.918 to 1.105, mean 1.040 (table B). 1.00 = schedule only", fill=YELLOW)
put(bk, "A%d" % r, "Implied equipment revenue at that ratio")
put(bk, "B%d" % r, "=B%d*B%d" % (BR["sched"], BR["uplift"]), BLACK, M)
BR["impl"] = r
r += 1
put(bk, "A%d" % r, "Implied vs Street")
put(bk, "B%d" % r, "=B%d/B%d-1" % (BR["impl"], BR["st"]), BLACK, PCT)
BR["impl_g"] = r
r += 1
bin_("drop", "Share of extra equipment revenue that reaches EBITDA", 0.15, PCT, "Equipment gross margin 10-13%; incremental 14-18%. Your assumption", fill=YELLOW)
put(bk, "A%d" % r, "Extra EBITDA over the four quarters ($M)")
put(bk, "B%d" % r, "=(B%d-B%d)*B%d" % (BR["impl"], BR["st"], BR["drop"]), BLACK, M)
BR["x_e"] = r
r += 2

header(bk, r, ["B. Back-test: scheduled vs actual next-four-quarter equipment revenue", "Backlog", "1-yr share", "Scheduled", "Actual", "Ratio", "Ex-Prolec", "Source"])
r += 1
BT = [("12/31/23", 40478, 0.51, 18952, 0.918, "Information statement (S-1/A 0001193125-24-131304)"),
      ("3/31/24", 42210, 0.45, 19532, 1.028, "10-Q 0001996810-24-000008"),
      ("6/30/24", 41561, 0.47, 20232, 1.036, "10-Q 0001996810-24-000067"),
      ("9/30/24", 42069, 0.46, 20822, 1.076, "10-Q 0001996810-24-000083"),
      ("12/31/24", 43047, 0.44, 20934, 1.105, "10-K 0001996810-25-000011"),
      ("3/31/25", 45478, 0.45, 21991, 1.052, "10-Q 0001996810-25-000073"),
      ("6/30/25", 49712, 0.42, 23556, 1.066, "10-Q 0001996810-25-000133")]
BR["bt_first"] = r
for d, rpo, sh, act, ex, src in BT:
    put(bk, "A%d" % r, d, BLUE)
    put(bk, "B%d" % r, rpo, BLUE, M)
    put(bk, "C%d" % r, sh, BLUE, PCT)
    put(bk, "D%d" % r, "=B%d*C%d" % (r, r), BLACK, M)
    put(bk, "E%d" % r, act, BLUE, M)
    put(bk, "F%d" % r, "=E%d/D%d" % (r, r), BLACK, '0.000')
    put(bk, "G%d" % r, ex, BLUE, '0.000')
    put(bk, "H%d" % r, src + ". Actual = equipment revenue in the next four quarters", NOTE)
    r += 1
BR["bt_last"] = r - 1
put(bk, "A%d" % r, "Mean (ex-Prolec)", BOLD)
put(bk, "G%d" % r, "=AVERAGE(G%d:G%d)" % (BR["bt_first"], BR["bt_last"]), BLACK, '0.000', bold=True)
BR["bt_mean"] = r
r += 1
put(bk, "A%d" % r, "Lowest / highest (ex-Prolec)")
put(bk, "F%d" % r, "=MIN(G%d:G%d)" % (BR["bt_first"], BR["bt_last"]), BLACK, '0.000')
put(bk, "G%d" % r, "=MAX(G%d:G%d)" % (BR["bt_first"], BR["bt_last"]), BLACK, '0.000')
put(bk, "H%d" % r, "Ex-Prolec deducts Prolec equipment revenue: 469 (Q1 2026) and 1,303 (H1 2026), 10-Q organic revenue table, 'Less: Acquisitions'. The 12/31/23 miss includes nuclear activities sold to EDF in 2024", NOTE)
r += 2

header(bk, r, ["C. The Street's Power equipment revenue in gigawatts and price per kW", "FY2026", "FY2027", "FY2028", "FY2029", "", "", "Source"])
r += 1
put(bk, "A%d" % r, "Street Power equipment revenue")
for col, sc in zip("BCDE", "CDEF"):
    put(bk, "%s%d" % (col, r), "=Street!%s%d" % (sc, SR["s_p_eq"]), GREEN, M)
BR["c_rev"] = r
r += 1
put(bk, "A%d" % r, "GW shipped (your Drivers)")
for col, dc in zip("BCDE", "CDEF"):
    put(bk, "%s%d" % (col, r), "=Drivers!%s%d" % (dc, DR["gw"]), GREEN, GW)
BR["c_gw"] = r
r += 1
put(bk, "A%d" % r, "Street revenue per kW at those GW")
for col in "BCDE":
    put(bk, "%s%d" % (col, r), "=IF(%s%d=0,0,%s%d/%s%d)" % (col, BR["c_gw"], col, BR["c_rev"], col, BR["c_gw"]), BLACK, KW)
BR["c_kw"] = r
r += 1
bin_("bk_price", "Signed Power equipment backlog per gas GW, 6/30/26 ($/kW, all-in)", 741, KW, "$39,261M Power equipment RPO (10-Q Q2 2026 MD&A p.29) / 53 GW Gas Power equipment backlog (Q2 2026 release 0001996810-26-000147). Numerator includes non-gas equipment (est. ~$2.6B; about $692/kW gas only)")
bin_("p_rpo", "Power equipment backlog, 6/30/26 ($M)", 39261, M, "10-Q Q2 2026 MD&A, Power")
put(bk, "A%d" % r, "Signed Power backlog (all delivery years) / Street Power equipment revenue, H2 2026 to 2028")
put(bk, "B%d" % r, "=B%d/(Street!C%d-History!F%d+Street!D%d+Street!E%d)" % (BR["p_rpo"], SR["s_p_eq"], HR["p_eq"], SR["s_p_eq"], SR["s_p_eq"]), BLACK, PCT)
put(bk, "H%d" % r, "Not a coverage ratio: part of the backlog delivers in 2029-30 (CEO: 2026 orders deliver mostly in 2029-30), and not all RPO converts on schedule. It shows the size of the signed book against the Street's revenue", NOTE)
BR["cover"] = r
r += 2

header(bk, r, ["D. Order prices and order-year margins", "Value", "", "", "", "", "", "Source"])
r += 1
for label, v, fmt, src in [
    ("Power equipment orders per gas GW ordered, Q4 2025 ($/kW, all-in)", 753, KW, "(Change in Power equipment RPO + Power equipment revenue) / gas GW ordered (10-K / 10-Q). Includes non-gas orders; moves with unit mix"),
    ("Power equipment orders per gas GW ordered, Q1 2026 ($/kW, all-in)", 705, KW, "Same method (10-Q Q1 2026: 8.1 GW). Management says H1 2026 orders were priced above Q4 2025; the dip is mix"),
    ("Power equipment orders per gas GW ordered, Q2 2026 ($/kW, all-in)", 1049, KW, "Same method; lifted by 61 aeroderivative units and combined-cycle scope (Q2 2026 call)"),
    ("Power order-year margin above the 2022 base: 2024 orders (points)", 7.5, '0.0', "Derived from disclosed changes in backlog margin and RPO (estimate, range 5.5-7.6)"),
    ("Power order-year margin above the 2022 base: 2025 orders (points)", 20.4, '0.0', "Derived (estimate, range 19.3-20.4)"),
    ("Power order-year margin above the 2022 base: 2026 orders (points)", 37.0, '0.0', "Derived (estimate, range 34-40); rests on one CEO statement (2026-09-16). CEO: 2026 orders deliver mostly in 2029-30"),
]:
    put(bk, "A%d" % r, label)
    put(bk, "B%d" % r, v, BLUE, fmt)
    put(bk, "H%d" % r, src, NOTE)
    r += 1
r += 1
header(bk, r, ["E. Gigawatts: output ramp and demand cover", "Q3 2025", "Q4 2025", "Q1 2026", "Q2 2026", "", "", "Source"])
r += 1
put(bk, "A%d" % r, "Gas turbine GW sold")
for col, v in zip("BCDE", (4.0, 3.1, 4.2, 3.3)):
    put(bk, "%s%d" % (col, r), v, BLUE, GW)
put(bk, "H%d" % r, "Power table, Sales in units, Gas Turbine Gigawatts: 10-Q Q3 2025 (4.0); 10-K FY2025 (15.3 for the year less 12.2 for nine months); 10-Q Q1 2026 (4.2); 10-Q Q2 2026 (3.3)", NOTE)
BR["e_gw"] = r
r += 1
put(bk, "A%d" % r, "GW sold, last four quarters", BOLD)
put(bk, "B%d" % r, "=SUM(B%d:E%d)" % (BR["e_gw"], BR["e_gw"]), BLACK, GW, bold=True)
BR["e_ltm"] = r
r += 1
bin_("e_rate", "Company output rate from Q3 2026 (GW a quarter)", 5, GW, "Q2 2026 release (0001996810-26-000147): 'on track to deliver 20 GW of annual gas turbine output in the third quarter of 2026'. CEO, 2026-09-16: 5 GW a quarter in Q3 and Q4 2026")
put(bk, "A%d" % r, "Step up needed against the last four quarters")
put(bk, "B%d" % r, "=B%d*4/B%d-1" % (BR["e_rate"], BR["e_ltm"]), BLACK, PCT)
put(bk, "H%d" % r, "The main execution test. GW sold in a quarter can differ from factory output: revenue on customised equipment is recognised over time", NOTE)
BR["e_step"] = r
r += 1
bin_("e_firm", "Firm gas backlog, 6/30/26 (GW)", 53, GW, "Q2 2026 release: 'backlog growth from 44 to 53 GW and an increase in slot reservation agreements from 56 to 63 GW'")
bin_("e_res", "Slot reservations, 6/30/26 (GW)", 63, GW, "Same release. A paid place in the factory queue: not yet an order and not in RPO. 10-K FY2025: reservations may not result in orders")
put(bk, "A%d" % r, "Signed or reserved (GW)", BOLD)
put(bk, "B%d" % r, "=B%d+B%d" % (BR["e_firm"], BR["e_res"]), BLACK, GW, bold=True)
BR["e_tot"] = r
r += 1
bin_("e_2030", "Company output goal for 2030 (GW)", 30, GW, "Q2 2026 release: 'we are implementing actions to produce 30 GW in 2030'")
put(bk, "A%d" % r, "Planned output, H2 2026 to 2030 (GW)")
put(bk, "B%d" % r, "=Drivers!C%d-History!F%d+Drivers!D%d+Drivers!E%d+Drivers!F%d+B%d" % (DR["gw"], HR["gw"], DR["gw"], DR["gw"], DR["gw"], BR["e_2030"]), GREEN, GW)
put(bk, "H%d" % r, "Your GW on Drivers for FY2026 (less H1 2026 actual) to FY2029, plus the 2030 goal", NOTE)
BR["e_plan"] = r
r += 1
put(bk, "A%d" % r, "Signed or reserved / planned output to 2030", BOLD)
put(bk, "B%d" % r, "=B%d/B%d" % (BR["e_tot"], BR["e_plan"]), BLACK, PCT, bold=True)
put(bk, "H%d" % r, "Some firm orders deliver in 2031, and reservations may not become orders. It shows that factory output, not new demand, is what is left to deliver the Street's Power equipment revenue", NOTE)
BR["e_cover"] = r
r += 2

header(bk, r, ["F. Price per kW the Street assumes, by half-year", "H1 2026 actual", "H2 2026", "H1 2027", "H2 2027", "", "", "Source"])
r += 1
put(bk, "A%d" % r, "Power equipment revenue ($M)")
put(bk, "B%d" % r, "=History!F%d" % HR["p_eq"], GREEN, M)
for col, k in zip("CDE", ("p_eq_h2_26", "p_eq_h1_27", "p_eq_h2_27")):
    put(bk, "%s%d" % (col, r), round(WIN[k]), BLUE, M)
put(bk, "H%d" % r, "H1 2026: 10-Q Q2 2026. Other columns: Bloomberg, 2026-10-01, quarterly division consensus (Total Revenue, Power, Equipment)", NOTE)
BR["h_rev"] = r
r += 1
put(bk, "A%d" % r, "Gas GW on the company's path")
put(bk, "B%d" % r, "=History!F%d" % HR["gw"], GREEN, GW)
put(bk, "C%d" % r, "=2*$B$%d" % BR["e_rate"], BLACK, GW)
put(bk, "D%d" % r, "=2*$B$%d" % BR["e_rate"], BLACK, GW)
put(bk, "E%d" % r, "=Drivers!D%d-D%d" % (DR["gw"], r), GREEN, GW)
put(bk, "H%d" % r, "H2 2026 and H1 2027 at the stated 5 GW a quarter (CEO, 2026-09-16: the start of 2027 will be 'very similar'). H2 2027 = your FY2027 GW on Drivers less H1 2027. These GW are the company's path, not a Street figure", NOTE)
BR["h_gw"] = r
r += 1
put(bk, "A%d" % r, "Revenue per kW ($/kW, all-in)", BOLD)
for col in "BCDE":
    put(bk, "%s%d" % (col, r), "=IF(%s%d=0,0,%s%d/%s%d)" % (col, BR["h_gw"], col, BR["h_rev"], col, BR["h_gw"]), BLACK, KW, bold=True)
put(bk, "H%d" % r, "The Street holds revenue per kW about flat from H2 2026 into H1 2027, then steps up. One half-year is a noisy test: reported revenue per kW was $365 in H1 2025 and $513 in H1 2026", NOTE)
BR["h_kw"] = r
r += 2

header(bk, r, ["G. HA turbines (largest class): which orders ship next", "Units", "", "", "", "", "", "Source"])
r += 1
bin_("ha_2412", "HA in backlog, 12/31/24", 32, '0', "10-K FY2024 (0001996810-25-000011): 'As of December 31, 2024, we had 32 HA-Turbines in RPO'")
bin_("ha_o25", "HA ordered in 2025", 43, '0', "10-K FY2025 Power table, Orders in units")
bin_("ha_s25", "HA sold in 2025", 24, '0', "10-K FY2025 Power table, Sales in units")
put(bk, "A%d" % r, "HA in backlog, 12/31/25")
put(bk, "B%d" % r, "=B%d+B%d-B%d" % (BR["ha_2412"], BR["ha_o25"], BR["ha_s25"]), BLACK, '0')
put(bk, "H%d" % r, "Ties to the 10-K FY2025: 'As of December 31, 2025, we had 51 HA-Turbines in RPO'", NOTE)
BR["ha_2512"] = r
r += 1
bin_("ha_o26", "HA ordered in H1 2026", 27, '0', "10-Q Q1 2026 and Q2 2026 Power tables, Orders in units (12 + 15)")
bin_("ha_s26", "HA sold in H1 2026", 8, '0', "10-Q Q1 2026 and Q2 2026 Power tables, Sales in units (5 + 3)")
put(bk, "A%d" % r, "HA in backlog, 6/30/26", BOLD)
put(bk, "B%d" % r, "=B%d+B%d-B%d" % (BR["ha_2512"], BR["ha_o26"], BR["ha_s26"]), BLACK, '0', bold=True)
put(bk, "H%d" % r, "Unit roll. The count itself is disclosed only in 10-Ks", NOTE)
BR["ha_bk"] = r
r += 1
put(bk, "A%d" % r, "Of which ordered in 2025 or later: at least")
put(bk, "B%d" % r, "=B%d-MAX(0,B%d-B%d)" % (BR["ha_bk"], BR["ha_2412"], BR["ha_s25"]), BLACK, '0')
put(bk, "H%d" % r, "Assumes no HA ordered in 2025 shipped in 2025 (HA lead times run well over a year). Then 8 older units were left at 12/31/25, and at worst all 8 are still in backlog", NOTE)
BR["ha_min"] = r
r += 1
put(bk, "A%d" % r, "Of which ordered in 2025 or later: at most")
put(bk, "B%d" % r, "=B%d-MAX(0,B%d-B%d-B%d)" % (BR["ha_bk"], BR["ha_2412"], BR["ha_s25"], BR["ha_s26"]), BLACK, '0')
put(bk, "H%d" % r, "If the oldest units shipped first: HA sold in 2025 and H1 2026 (32) equal the 12/31/24 backlog (32), so every older unit has shipped. Order year of shipped units is not disclosed. HA only: other turbine classes may still carry older orders", NOTE)
BR["ha_max"] = r
r += 1
widths(bk, {"A": 66, "B": 13, "C": 11, "D": 11, "E": 11, "F": 9, "G": 9, "H": 100})
bk.freeze_panes = "B4"

# ======================================================================= Summary
sm = wb.create_sheet("Summary", 0)
title(sm, "GE Vernova (GEV): model summary", "All values link to the other sheets. Direction and judgments are set on Drivers and Valuation.")
r = 4
header(sm, r, ["Market", "Value", "", "", ""])
r += 1
SMR = {}


def sline(label, f, fmt, key, bold=False):
    global r
    put(sm, "A%d" % r, label, BOLD if bold else BLACK)
    put(sm, "B%d" % r, f, GREEN, fmt, bold=bold)
    SMR[key] = r
    r += 1


sline("Share price ($), Bloomberg 2026-10-01", "=Valuation!B%d" % VR["px"], PX, "px")
sline("Market capitalisation ($M)", "=Valuation!B%d" % VR["mcap"], M, "mcap")
sline("Enterprise value ($M)", "=Valuation!B%d" % VR["ev"], M, "ev")
sline("Forward EV/EBITDA today (Bloomberg)", "=Valuation!B%d" % VR["m_bbg"], X, "mbbg")
r += 1
header(sm, r, ["Price in about 12 months (30 Sep 2027)", "Bear: GEV plan path", "Base: your forecast", "Bull: base + uplift", ""])
r += 1
for label, key, fmt in [("Next-twelve-month EBITDA ($M)", "ntm", M), ("Exit multiple", "mult", X), ("Price per share ($)", "sp", PX), ("Return", "sr", PCT), ("Probability", "prob", PCT)]:
    put(sm, "A%d" % r, label, BOLD if key in ("sp", "sr") else BLACK)
    for col in "BCD":
        put(sm, "%s%d" % (col, r), "=Valuation!%s%d" % (col, VR[key]), GREEN, fmt, bold=key in ("sp", "sr"))
    SMR[key] = r
    r += 1
sline("Probability-weighted price ($)", "=Valuation!B%d" % VR["pw"], PX, "pw", bold=True)
sline("Probability-weighted return", "=Valuation!B%d" % VR["pwr"], PCT, "pwr", bold=True)
sline("Break-even exit multiple (base)", "=Valuation!B%d" % VR["be"], X, "be")
sline("Base return from EBITDA growth and cash added (multiple unchanged)", "=Valuation!B%d" % VR["r_growth"], PCT, "rg")
sline("Base return from the change in multiple", "=Valuation!B%d" % VR["r_mult"], PCT, "rm")
sline("Base EBITDA at the target date vs Street", "=Valuation!B%d" % VR["vs_st"], PCT, "vsst")
put(sm, "C%d" % SMR["rg"], "Split uses today's Bloomberg multiple (28.5x)", NOTE)
put(sm, "C%d" % SMR["vsst"], "Near 0% with default drivers: the base case is the Street's earnings; the return depends on the multiple", NOTE)
r += 1
header(sm, r, ["Adjusted EBITDA ($M)", "FY2026E", "FY2027E", "FY2028E", "FY2029E"])
r += 1
for label, row, sheet in [("Model", FR["ebitda"], "Forecast"), ("Street (Bloomberg annual mean)", FR["st_e"], "Forecast"), ("Model vs Street (defaults: segment vs annual consensus)", FR["vs_e"], "Forecast")]:
    put(sm, "A%d" % r, label)
    for col, fcol in zip("BCDE", "CDEF"):
        put(sm, "%s%d" % (col, r), "=%s!%s%d" % (sheet, fcol, row), GREEN, PCT if "vs" in label else M)
    r += 1
put(sm, "A%d" % r, "Model adjusted EBITDA margin")
for col, fcol in zip("BCDE", "CDEF"):
    put(sm, "%s%d" % (col, r), "=Forecast!%s%d" % (fcol, FR["m"]), GREEN, PCT)
r += 2
header(sm, r, ["Backlog and gigawatt checks", "Value", "", "", ""])
r += 1
sline("Equipment revenue scheduled from backlog, Q3 2026-Q2 2027 ($M)", "=Backlog!B%d" % BR["sched"], M, "sched")
sline("Street equipment revenue, same quarters ($M)", "=Backlog!B%d" % BR["st"], M, "st")
sline("Schedule above the Street", "=Backlog!B%d" % BR["gap"], PCT, "gap")
sline("Signed Power backlog (all years) / Street Power equipment revenue H2 2026-2028", "=Backlog!B%d" % BR["cover"], PCT, "cover")
put(sm, "C%d" % SMR["cover"], "Size of the signed book, not coverage: part delivers 2029-30", NOTE)
sline("Gas GW sold, last four quarters", "=Backlog!B%d" % BR["e_ltm"], GW, "ltm")
sline("Step up needed to reach the stated 5 GW a quarter", "=Backlog!B%d" % BR["e_step"], PCT, "step")
sline("Signed or reserved GW / planned output to 2030", "=Backlog!B%d" % BR["e_cover"], PCT, "gwcover")
sline("HA turbines in backlog, 6/30/26 (units)", "=Backlog!B%d" % BR["ha_bk"], '0', "habk")
sline("Of which ordered in 2025 or later, at least (units)", "=Backlog!B%d" % BR["ha_min"], '0', "hamin")
r += 1
header(sm, r, ["Checks", "Status", "", "", ""])
r += 1
put(sm, "A%d" % r, "History ties to reported totals (max absolute difference)")
put(sm, "B%d" % r, "=MAX(ABS(MAX(History!B%d:F%d)),ABS(MIN(History!B%d:F%d)),ABS(MAX(History!B%d:F%d)),ABS(MIN(History!B%d:F%d)),ABS(MAX(History!B%d:F%d)),ABS(MIN(History!B%d:F%d)))" % (
    HR["chk_rev"], HR["chk_rev"], HR["chk_rev"], HR["chk_rev"], HR["chk_seg"], HR["chk_seg"], HR["chk_seg"], HR["chk_seg"], HR["chk_e"], HR["chk_e"], HR["chk_e"], HR["chk_e"]), GREEN, M1)
SMR["chk_h"] = r
r += 1
put(sm, "A%d" % r, "FY2026 revenue vs guide")
put(sm, "B%d" % r, "=Forecast!C%d" % FR["chk_g1"], GREEN)
r += 1
put(sm, "A%d" % r, "FY2026 margin vs guide")
put(sm, "B%d" % r, "=Forecast!C%d" % FR["chk_g2"], GREEN)
r += 1
put(sm, "A%d" % r, "Scenario probabilities sum to 100%")
put(sm, "B%d" % r, "=Valuation!B%d" % VR["pchk"], GREEN, PCT)
r += 2
put(sm, "A%d" % r, "Colour key: blue = typed input; yellow = your assumption; black = formula; green = link to another sheet.", NOTE)
r += 2
header(sm, r, ["How to use this model", "", "", "", ""])
r += 1
for line in [
    "1. Forecast: change the yellow cells on Drivers (GW, price per kW, growth, margins, Wind, corporate). Forecast recalculates.",
    "2. Valuation: change the yellow cells on Valuation (exit multiples, probabilities, bull uplift, target date, cash added).",
    "3. Backlog: change the ratio and drop-through on Backlog to test the equipment-schedule check.",
    "4. Do not type over black or green cells; they are formulas.",
    "5. Before submitting: File > Info > Properties, and remove any name from Author; the memo and model must be blinded.",
]:
    put(sm, "A%d" % r, line)
    r += 1
widths(sm, {"A": 60, "B": 14, "C": 14, "D": 14, "E": 14})

# ======================================================================= Sources
so = wb.create_sheet("Sources")
title(so, "Sources", "Accession numbers are SEC EDGAR filings by GE Vernova (CIK 1996810) unless noted.")
header(so, 4, ["Source", "Used for"])
rows = [
    ("10-K FY2025, 0001996810-26-000015 (MD&A p.26; Note 24 pp.72-73; income statement p.39; RPO note p.53)", "History FY2023-FY2025, segment data, backlog schedule at 12/31/25"),
    ("10-K FY2024, 0001996810-25-000011", "FY2023 equipment and services revenue; backlog schedule at 12/31/24; HA units in RPO at 12/31/24"),
    ("10-Q Q2 2026, 0001996810-26-000148 (income statement p.6; segment note p.23; MD&A p.29; RPO note; Note 14; cover)", "H1 2025 restated and H1 2026; bridge; shares; backlog schedule; Power backlog and GW"),
    ("10-Q Q1 2026, 0001996810-26-000064; earlier 10-Qs 0001996810-24-000008, -24-000067, -24-000083, -25-000073, -25-000133, -25-000160", "Back-test windows; order prices; unit tables"),
    ("Information statement S-1/A, 0001193125-24-131304", "Backlog schedule at 12/31/23"),
    ("Q2 2026 earnings release, 0001996810-26-000147 (2026-07-22)", "FY2026 guidance; gas backlog 53 GW and slot reservations 63 GW; output plan (20 GW rate in Q3 2026, 24 GW in 2028, 30 GW in 2030)"),
    ("Earnings releases 0001996810-25-000004 (FY2024) and 0001996810-26-000012 (FY2025)", "Reported adjusted EBITDA cross-check"),
    ("Q4 2025 earnings release, 0001996810-26-000012 (2026-01-28)", "2028 framework"),
    ("Morgan Stanley Laguna conference transcript, gevernova.com, 2026-09-16", "Output path; timing of order margins"),
    ("Bloomberg, 2026-10-01 (BEst consensus; EEO company-specific quarterly consensus; Snapshot; Multiples)", "Street sheet; Backlog table A; Valuation market data"),
]
for i, (s, u) in enumerate(rows, 5):
    put(so, "A%d" % i, s)
    put(so, "B%d" % i, u)
widths(so, {"A": 110, "B": 80})

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("wrote", OUT)
