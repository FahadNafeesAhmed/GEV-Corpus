"""Builds the team thesis file as a plain black-and-white PDF with side notes (research notes, not submission text)."""
import datetime as dt
import re
import subprocess
import sys
import warnings
from pathlib import Path

import markdown
from openpyxl import load_workbook

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
import dossier_charts as ch  # noqa: E402
import content_a, content_b, content_c  # noqa: E402

warnings.filterwarnings("ignore")
OUT_DIR = HERE.parent.parent / "docs" / "research" / "multi_agent" / "14_full_thesis"
PULL1 = str(HERE.parent.parent / "data" / "bloomberg" / "GEV_bloomberg_pull_filled_2026-10-01.xlsx")
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
DARK, MID, LIGHT = "#1a1a1a", "#8c8c8c", "#cfcfcf"
ch.BLUE, ch.INK, ch.GRID = DARK, "#000000", "#dddddd"


def usd(v):
    return "$" + format(int(round(v)), ",")


def num(v):
    return format(int(round(v)), ",")


def series(ws, c0):
    out = []
    for d, v in ws.iter_rows(min_row=4, min_col=c0, max_col=c0 + 1, values_only=True):
        if isinstance(d, (dt.datetime, dt.date)) and isinstance(v, (int, float)):
            out.append((d.date() if isinstance(d, dt.datetime) else d, float(v)))
    return out


def charts():
    wb = load_workbook(PULL1, data_only=True, read_only=True)
    mult = series(wb["Multiples"], 1)
    px = series(wb["PricePositioning"], 1)
    hi = max(px, key=lambda x: x[1])
    print("price series %s to %s, n=%d, high %.2f on %s | multiple n=%d last %.2f" % (px[0][0], px[-1][0], len(px), hi[1], hi[0], len(mult), mult[-1][1]))
    C = {}
    ch.W = 540
    C["cases"] = ch.bars("Share price in three cases", ["Bear", "Today", "Base", "Bull"],
                         [dict(vals=[855, None, 1126, 1341], color=DARK, txt=["$855", "", "$1,126", "$1,341"]),
                          dict(vals=[None, 987.45, None, None], color=MID, txt=["", "$987", "", ""])], 1500, usd, stacked=True, maxbar=70, h=230)
    C["bridge"] = ch.waterfall("Price bridge", ["Today", "Profit growth", "Lower multiple", "Base target"],
                               [(0, 987.45, MID, "$987"), (987.45, 1327.82, DARK, "+$340"), (1126.02, 1327.82, LIGHT, "-$202"), (0, 1126.02, MID, "$1,126")], 1500, usd, h=240)
    C["ebitda"] = ch.bars("Adjusted EBITDA", ["2023", "2024", "2025", "2026E", "2027E", "2028E", "2029E"],
                          [dict(vals=[807, 2035, 3196, None, None, None, None], color=DARK, txt=["807", "2,035", "3,196", "", "", "", ""]),
                           dict(vals=[None, None, None, 6334, 9471, 12553, 16496], color=DARK, outlined=True, txt=["", "", "", "6,334", "9,471", "12,553", "16,496"])],
                          18000, num, stacked=True, maxbar=44, nt=6, h=250)
    C["stack"] = ch.bars("Gas backlog and reservations", ["Jun 2025", "Sep 2025", "Dec 2025", "Mar 2026", "Jun 2026"],
                         [dict(vals=[29, 33, 40, 44, 53], color=DARK, txt=["29", "33", "40", "44", "53"]),
                          dict(vals=[25, 29, 43, 56, 63], color=MID, txt=["25", "29", "43", "56", "63"])], 125, lambda v: "%g" % v, stacked=True, maxbar=60, h=250)
    q = ["Q1 25", "Q2 25", "Q3 25", "Q4 25", "Q1 26", "Q2 26", "Q3 26", "Q4 26", "Q1 27", "Q2 27"]
    C["qship"] = ch.bars("Gigawatts shipped each quarter", q,
                         [dict(vals=[3.0, 5.2, 4.0, 3.1, 4.2, 3.3] + [None] * 4, color=DARK, txt=["3.0", "5.2", "4.0", "3.1", "4.2", "3.3", "", "", "", ""]),
                          dict(vals=[None] * 6 + [5, 5, 5, 5], color=DARK, outlined=True, txt=[""] * 6 + ["5", "5", "5", "5"])], 6, lambda v: "%g" % v, stacked=True, maxbar=34, nt=6, h=240)
    C["kw"] = ch.lines("Revenue per kW", 8,
                       [dict(vals=[406, 480, 437, 513, None, None, None, None], color=DARK, marker=True, txt=["406", "480", "437", "513", "", "", "", ""]),
                        dict(vals=[None, None, None, None, 534, 579, 662, 741], color=MID, dash="6 4", marker=True, txt=["", "", "", "", "534", "579", "662", "741"])],
                       300, 900, usd, list(enumerate(["2023", "2024", "2025", "H1 26", "2026E", "2027E", "2028E", "2029E"])), hlines=[(741, "Signed backlog: $741 per kW", DARK)], nt=6, pad=26, h=250)
    C["bt"] = ch.bars("Backlog schedule back-test", ["Dec 23", "Mar 24", "Jun 24", "Sep 24", "Dec 24", "Mar 25", "Jun 25"],
                      [dict(vals=[0.918, 1.028, 1.036, 1.076, 1.105, 1.052, 1.066], color=DARK, txt=["0.92", "1.03", "1.04", "1.08", "1.11", "1.05", "1.07"])],
                      1.2, lambda v: "%.2f" % v, ymin=0.8, hline=(1.0, "", DARK), maxbar=40, nt=4, h=220)
    C["margin"] = ch.bars("Power margin", ["2023", "2024", "2025", "H1 26", "2026E", "2027E", "2028E", "2029E"],
                          [dict(vals=[9.9, 12.5, 14.7, 17.6, None, None, None, None], color=DARK, txt=["9.9", "12.5", "14.7", "17.6", "", "", "", ""]),
                           dict(vals=[None, None, None, None, 18.7, 21.2, 23.3, 25.6], color=DARK, outlined=True, txt=["", "", "", "", "18.7", "21.2", "23.3", "25.6"])],
                          30, lambda v: "%g" % v, stacked=True, hline=(22, "Company target for 2028: 22%", DARK), maxbar=40, h=240)
    C["howmet"] = ch.bars("Howmet revenue growth", ["Q1 25", "Q2 25", "Q3 25", "Q4 25", "Q1 26", "Q2 26"],
                          [dict(vals=[6.5, 9.2, 13.8, 14.6, 19.1, 24.1], color=DARK, txt=["6.5", "9.2", "13.8", "14.6", "19.1", "24.1"])], 45, lambda v: "%g" % v,
                          hline=(38, "Gas turbine market alone, Q2 2026: +38%", DARK), maxbar=44, nt=3, h=230)
    xt = [(i, d.strftime("%b %y")) for i, (d, v) in enumerate(mult) if i % 24 == 0 and i < len(mult) - 8]
    C["mult"] = ch.lines("Forward EV/EBITDA", len(mult), [dict(vals=[v for d, v in mult], color=DARK)], 10, 45, lambda v: "%gx" % v, xt,
                         hlines=[(25, "Base exit multiple: 25x", DARK), (21.8, "Break-even: 21.8x", DARK)], h=250, nt=7)
    pxt = [(i, d.strftime("%b %y")) for i, (d, v) in enumerate(px) if i % 170 == 0 and i < len(px) - 60]
    sell_i = min(range(len(px)), key=lambda i: abs((px[i][0] - dt.date(2026, 9, 14)).days))
    C["price"] = ch.lines("Share price", len(px), [dict(vals=[v for d, v in px], color=DARK)], 0, 1500, usd, pxt, vlines=[(sell_i, "Sell note, 14 Sep 2026", DARK)],
                          hlines=[(1126, "Base target: $1,126", DARK)], h=250, nt=5)
    peers = [("GE Vernova (+50%)", 28.5), ("Doosan Enerbility (+28%)", 28.2), ("Howmet (+18%)", 25.7), ("GE Aerospace (+13%)", 25.0), ("ATI (+18%)", 20.6), ("Woodward (+11%)", 20.0),
             ("Caterpillar (+19%)", 19.3), ("Carpenter (+14%)", 17.7), ("Mitsubishi Heavy (+17%)", 15.1), ("Siemens Energy (+31%)", 12.7), ("Baker Hughes (+17%)", 11.1)]
    C["peers"] = ch.hbars("Multiples and growth", [p[0] for p in peers], [p[1] for p in peers], [DARK] + [MID] * 10, lambda v: "%.1fx" % v, 34)
    be = [29.61, 28.75, 27.93, 27.16, 26.44, 25.76, 25.12, 24.50, 23.92, 23.36, 22.83, 22.32, 21.83]
    path = [29.61 + (25.0 - 29.61) * k / 12.0 for k in range(13)]
    C["horizon"] = ch.lines("Hurdle by month", 13,
                            [dict(vals=be, color=DARK, marker=True, txt=["29.6x", "", "", "27.2x", "", "", "25.1x", "", "", "23.4x", "", "", "21.8x"]),
                             dict(vals=path, color=MID, dash="6 4")], 20, 31, lambda v: "%gx" % v,
                            [(0, "Today"), (3, "3 months"), (6, "6 months"), (9, "9 months"), (12, "12 months")], h=250, nt=11, pad=30)
    cats = ["Dec 23", "Mar 24", "Jun 24", "Sep 24", "Dec 24", "Mar 25", "Jun 25", "Jun 26"]
    C["elec"] = ch.bars("Electrification conversion", cats,
                        [dict(vals=[46.2, 43.0, 39.6, 40.4, 41.1, 40.2, 39.3, None], color=DARK, txt=["46.2", "43.0", "39.6", "40.4", "41.1", "40.2", "39.3", ""]),
                         dict(vals=[None] * 7 + [36.2], color=DARK, outlined=True, txt=[""] * 7 + ["36.2"])], 50, lambda v: "%g" % v, ymin=30, stacked=True, maxbar=40, nt=4, h=220)
    C["flow"] = flow()
    return {k: v.replace("Arial, Helvetica, sans-serif", "'Times New Roman', Times, serif") for k, v in C.items()}


def flow():
    boxes = [("1  Corpus", "filings, calls, news"), ("2  Factory physics", "blades, foundries"), ("3  Past winners", "20 pitches read"),
             ("4  Bloomberg", "three data pulls"), ("5  Agent tests", "over 400 runs"), ("6  Stress tests", "25 tests, both sides")]
    w, h, gap, x0, y0 = 108, 46, 18, 4, 8
    out = ['<svg viewBox="0 0 760 122" width="100%" xmlns="http://www.w3.org/2000/svg" font-family="\'Times New Roman\', Times, serif">']
    for i, (a, b) in enumerate(boxes):
        x = x0 + i * (w + gap)
        out.append('<rect x="%d" y="%d" width="%d" height="%d" fill="#fff" stroke="#000" stroke-width="1.2"/>' % (x, y0, w, h))
        out.append('<text x="%d" y="%d" font-size="13" font-weight="bold" text-anchor="middle">%s</text>' % (x + w / 2, y0 + 19, a))
        out.append('<text x="%d" y="%d" font-size="11.5" text-anchor="middle">%s</text>' % (x + w / 2, y0 + 36, b))
        if i < len(boxes) - 1:
            xa = x + w
            out.append('<line x1="%d" y1="%d" x2="%d" y2="%d" stroke="#000" stroke-width="1.2"/>' % (xa + 2, y0 + h / 2, xa + gap - 5, y0 + h / 2))
            out.append('<polygon points="%d,%d %d,%d %d,%d" fill="#000"/>' % (xa + gap - 1, y0 + h / 2, xa + gap - 7, y0 + h / 2 - 4, xa + gap - 7, y0 + h / 2 + 4))
    out.append('<line x1="380" y1="%d" x2="380" y2="76" stroke="#000" stroke-width="1.2"/>' % (y0 + h))
    out.append('<line x1="58" y1="%d" x2="702" y2="%d" stroke="#000" stroke-width="0.8" stroke-dasharray="3 3"/>' % (y0 + h + 9, y0 + h + 9))
    out.append('<rect x="190" y="78" width="380" height="36" fill="#1a1a1a"/>')
    out.append('<text x="380" y="101" font-size="13.5" font-weight="bold" fill="#fff" text-anchor="middle">Long, 12 months: base $1,126, break-even 21.8x</text>')
    out.append("</svg>")
    return "".join(out)


CSS = """
@page { size: Letter; margin: 0.6in 0.55in 0.7in 0.55in;
  @bottom-right { content: counter(page) " / " counter(pages); font: 8pt 'Times New Roman', serif; }
  @bottom-left { content: "GE Vernova: thesis file for the team. Research notes, not submission text."; font: 8pt 'Times New Roman', serif; } }
* { box-sizing: border-box; }
body { font-family: 'Times New Roman', Times, serif; font-size: 10pt; line-height: 1.32; color: #000; margin: 0; }
.col { width: 65.5%; }
h1 { font-size: 21pt; margin: 0 0 2pt; line-height: 1.1; width: 152.7%; }
.sub { font-size: 10pt; margin: 0 0 6pt; width: 152.7%; }
.box { border: 1pt solid #000; padding: 5pt 8pt; font-size: 9pt; margin: 6pt 0 8pt; width: 152.7%; }
.toc { font-size: 9pt; columns: 2; margin: 0 0 4pt; padding: 0; list-style: none; }
h2 { font-size: 13pt; border-bottom: 1.5pt solid #000; padding-bottom: 2pt; margin: 15pt 0 6pt; break-after: avoid; clear: both; width: 152.7%; }
h3 { font-size: 10.5pt; margin: 9pt 0 3pt; break-after: avoid; }
p { margin: 0 0 5pt; }
ul, ol { margin: 0 0 6pt; padding-left: 15pt; }
li { margin: 0 0 3pt; }
table { border-collapse: collapse; width: 100%; font-size: 8.6pt; line-height: 1.22; margin: 3pt 0 8pt; }
th { text-align: left; border-top: 1pt solid #000; border-bottom: 0.75pt solid #000; padding: 2pt 4pt; vertical-align: bottom; }
td { border-bottom: 0.4pt solid #9a9a9a; padding: 2pt 4pt; vertical-align: top; }
tr { break-inside: avoid; }
tr:last-child td { border-bottom: 1pt solid #000; }
.wide { width: 152.7%; clear: both; }
aside.note { float: right; clear: right; width: 45.5%; margin: 2pt -52.7% 8pt 0; font-size: 8.5pt; line-height: 1.27; border-left: 1.5pt solid #000; padding: 0 0 0 7pt; break-inside: avoid; }
aside.note p { margin: 0; }
figure { margin: 5pt 0 9pt; break-inside: avoid; }
figcaption { font-size: 9pt; font-weight: bold; margin-bottom: 2pt; }
.src { font-size: 7.8pt; color: #222; margin-top: 1pt; }
a { color: #000; text-decoration: none; }
.refs { width: 152.7%; font-size: 8.2pt; line-height: 1.25; }
.refs p { margin: 0 0 3pt; padding-left: 26pt; text-indent: -26pt; word-break: break-word; }
.refs a { text-decoration: underline; }
code { font-family: 'Times New Roman', serif; font-style: italic; }
"""


def md(text):
    return markdown.markdown(text.strip(), extensions=["tables"])


def link_refs(html):
    return re.sub(r"\[R(\d+)\]", r'<a href="#R\1">[R\1]</a>', html)


def render(sections, C):
    out = []
    for title, blocks in sections:
        out.append("<h2>%s</h2>" % title)
        for b in blocks:
            kind = b[0]
            if kind == "note":
                out.append('<aside class="note">%s</aside>' % md(b[1]))
            elif kind == "md":
                out.append(md(b[1]))
            elif kind == "wide":
                out.append('<div class="wide">%s</div>' % md(b[1]))
            elif kind in ("chart", "widechart"):
                src = ('<div class="src">Source: %s.</div>' % b[3]) if b[3] else ""
                out.append('<figure class="%s"><figcaption>%s</figcaption>%s%s</figure>' % ("wide" if kind == "widechart" else "", b[2], C[b[1]], src))
    return link_refs("\n".join(out))


def main():
    C = charts()
    sections = content_a.S + content_b.S + content_c.S
    titles = [t for t, _ in sections] + ["16. References", content_c.S_END[0][0]]
    refs = ['<h2>16. References</h2><div class="refs">']
    for key, text, url in content_c.REFS:
        link = (' <a href="%s">%s</a>' % (url, url)) if url else ""
        refs.append('<p id="%s"><b>%s</b>&nbsp;&nbsp; %s.%s</p>' % (key, key, text, link))
    refs.append("</div>")
    head = """<h1>GE Vernova (GEV): the thesis file for the team</h1>
<p class="sub">Long, 12 months, to 30 September 2027. Working date 2 October 2026. Money in $ millions unless stated.</p>
<div class="box"><b>Read this first.</b> These are research notes made with AI assistance. Under rule 10 of the competition that use must be disclosed, and the memo and the Excel model must be the team's own work: this file is not text to submit. Bloomberg figures are licensed for academic use. Keep this file inside the team and cite them as "Bloomberg" with the date.</div>
<aside class="note"><p><b>How to read this file.</b> The wide column holds the facts, tables and graphs. This narrow column says the same thing in plain words. Tags like [R3] point to the reference list in section 16.</p></aside>
<ul class="toc">%s</ul>""" % "".join("<li>%s</li>" % t for t in titles)
    body = head + render(sections, C) + "\n".join(refs) + render(content_c.S_END, C)
    html = '<!doctype html><html lang="en"><head><meta charset="utf-8"><title>GE Vernova thesis file</title><style>%s</style></head><body><div class="col">%s</div></body></html>' % (CSS, body)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    src = HERE / "GEV_FULL_THESIS.html"
    src.write_text(html, encoding="utf-8")
    pdf = OUT_DIR / "GEV_FULL_THESIS.pdf"
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--print-to-pdf=%s" % pdf, src.as_uri()], check=True, timeout=120,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print("html %d chars | pdf %s %d bytes" % (len(html), pdf, pdf.stat().st_size))


if __name__ == "__main__":
    main()
